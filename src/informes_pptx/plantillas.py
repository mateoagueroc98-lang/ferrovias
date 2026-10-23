"""Informes PowerPoint – plantillas.

Una plantilla sale de un PowerPoint de ejemplo. Se guarda en informes/plantillas/<id>/:
- base.pptx: copia del PowerPoint, con el identificador de cada lugar de imagen escrito como
  nombre del shape (así se encuentra después aunque cambie el orden).
- plantilla.json: nombre, fechas, tamaño de diapositiva y, por diapositiva, los textos (fijo o
  variable) y los lugares de imagen (nombre legible, posición y tamaño).
- miniaturas/ (opcional): una imagen por diapositiva, si se pudieron generar con PowerPoint.

Todo con python-pptx: no necesita PowerPoint instalado.
Para ver qué detecta en un PowerPoint:  python -m src.informes_pptx.plantillas archivo.pptx
"""
import datetime as dt
import io
import json
import re
import secrets
import shutil
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from src.informes_pptx import reemplazo

RAIZ = Path(__file__).resolve().parents[2] / "informes"
VERSION = 1
EMU_POR_CM = 360000
FIJO, VARIABLE = "fijo", "variable"
TODO, FRAGMENTO = "todo", "fragmento"

NOMBRES_MARCADOR = {
    PP_PLACEHOLDER.TITLE: "Título", PP_PLACEHOLDER.CENTER_TITLE: "Título",
    PP_PLACEHOLDER.SUBTITLE: "Subtítulo", PP_PLACEHOLDER.BODY: "Marcador de texto",
    PP_PLACEHOLDER.OBJECT: "Marcador de contenido", PP_PLACEHOLDER.FOOTER: "Pie de página",
    PP_PLACEHOLDER.DATE: "Fecha", PP_PLACEHOLDER.SLIDE_NUMBER: "N° de diapositiva",
}
# Marcadores vacíos que esperan una imagen o un gráfico: se toman como lugares de imagen
MARCADORES_IMAGEN = {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.BITMAP}
URI_OLE = "http://schemas.openxmlformats.org/presentationml/2006/ole"
URI_GRAFICO = "http://schemas.openxmlformats.org/drawingml/2006/chart"
# Elementos de la diapositiva que python-pptx no recorre (se avisa que hay algo no reconocido)
DESCONOCIDOS = {"{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent": "objeto especial",
                qn("p:contentPart"): "dibujo a mano"}


class ErrorPlantilla(ValueError):
    """El archivo o la plantilla tienen un problema (mensaje listo para mostrar)."""


# ------------------------------------------------------------------- utilidades
def slug(texto: str) -> str:
    """'Mes del título' -> 'mes_del_titulo' (para identificadores)."""
    t = unicodedata.normalize("NFD", str(texto).lower())
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def ahora() -> str:
    return dt.datetime.now().replace(microsecond=0).isoformat()


def cm(emu: int) -> float:
    return round(emu / EMU_POR_CM, 2)


@dataclass(frozen=True)
class _Transformacion:
    """Pasa coordenadas de adentro de un grupo a coordenadas de la diapositiva: X = dx + ex·x."""
    dx: float = 0.0
    dy: float = 0.0
    ex: float = 1.0
    ey: float = 1.0

    def caja(self, x, y, ancho, alto) -> dict:
        x, y, ancho, alto = (v or 0 for v in (x, y, ancho, alto))
        return {"x": round(self.dx + self.ex * x), "y": round(self.dy + self.ey * y),
                "ancho": round(self.ex * ancho), "alto": round(self.ey * alto)}

    def dentro_de(self, grupo) -> "_Transformacion":
        xfrm = grupo._element.grpSpPr.find(qn("a:xfrm"))
        partes = [xfrm.find(qn(f"a:{t}")) if xfrm is not None else None
                  for t in ("off", "ext", "chOff", "chExt")]
        if any(p is None for p in partes):
            return self
        off, ext, ch_off, ch_ext = partes
        ex = int(ext.get("cx")) / int(ch_ext.get("cx")) if int(ch_ext.get("cx")) else 1.0
        ey = int(ext.get("cy")) / int(ch_ext.get("cy")) if int(ch_ext.get("cy")) else 1.0
        dx = int(off.get("x")) - int(ch_off.get("x")) * ex
        dy = int(off.get("y")) - int(ch_off.get("y")) * ey
        return _Transformacion(self.dx + self.ex * dx, self.dy + self.ey * dy, self.ex * ex, self.ey * ey)


def _recorrer(shapes, transf=_Transformacion(), en_grupo=False):
    """Todos los shapes, entrando en los grupos: (shape, transformación, ¿está en un grupo?)."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _recorrer(shape.shapes, transf.dentro_de(shape), True)
        else:
            yield shape, transf, en_grupo


def _uri(shape) -> str:
    uris = shape._element.xpath("./a:graphic/a:graphicData/@uri")
    return uris[0] if uris else ""


def _tipo_marcador(shape):
    return shape.placeholder_format.type if shape.is_placeholder else None


def asegurar_ids_unicos(prs) -> int:
    """Cada shape tiene un número (id) que tiene que ser único en su diapositiva: los repetidos
    se renumeran. Devuelve cuántos cambió (casi siempre 0)."""
    cambiados = 0
    for slide in prs.slides:
        elementos = slide.shapes._spTree.xpath(".//p:cNvPr")
        maximo = max((int(e.get("id", 0)) for e in elementos), default=0)
        vistos = set()
        for e in elementos:
            if e.get("id") in vistos:
                maximo += 1
                e.set("id", str(maximo))
                cambiados += 1
            vistos.add(e.get("id"))
    return cambiados


# -------------------------------------------------------------------- detección
def _texto(numero, shape, transf, en_grupo, text_frame, origen, celda=None, caja=None) -> dict:
    sufijo = f"_f{celda[0]}c{celda[1]}" if celda else ""
    return {
        "id": f"t{numero:02d}_{shape.shape_id}{sufijo}",
        "shape_id": shape.shape_id,
        "celda": list(celda) if celda else None,
        "shape_nombre": shape.name,
        "origen": origen + (" (en un grupo)" if en_grupo else ""),
        "contenido": reemplazo.texto_cuadro(text_frame),
        "caja": caja or transf.caja(shape.left, shape.top, shape.width, shape.height),
        "tipo": FIJO,
        "reemplazo": TODO,
        "fragmento": "",
        "variable_id": "",
        "variable_nombre": "",
    }


def _celdas(numero, shape, transf, en_grupo):
    """Textos de las celdas de una tabla (las vacías no se listan)."""
    tabla = shape.table
    anchos = [c.width for c in tabla.columns]
    altos = [f.height for f in tabla.rows]
    for r, fila in enumerate(tabla.rows):
        for c, celda in enumerate(fila.cells):
            if celda.is_spanned or not reemplazo.texto_cuadro(celda.text_frame).strip():
                continue
            filas_unidas = celda.span_height if celda.is_merge_origin else 1
            cols_unidas = celda.span_width if celda.is_merge_origin else 1
            caja = transf.caja((shape.left or 0) + sum(anchos[:c]), (shape.top or 0) + sum(altos[:r]),
                               sum(anchos[c:c + cols_unidas]), sum(altos[r:r + filas_unidas]))
            yield _texto(numero, shape, transf, en_grupo, celda.text_frame,
                         f"Tabla «{shape.name}» (fila {r + 1}, columna {c + 1})", (r + 1, c + 1), caja)


def analizar(prs) -> dict:
    """Recorre el PowerPoint y arma la descripción de cada diapositiva (sin id ni nombre)."""
    diapositivas = []
    for numero, slide in enumerate(prs.slides, 1):
        textos, imagenes, avisos = [], [], []
        for shape, transf, en_grupo in _recorrer(slide.shapes):
            marcador = _tipo_marcador(shape)
            clase = None
            if isinstance(shape, Picture):
                clase = "Imagen"
            elif isinstance(shape, GraphicFrame):
                if shape.has_chart or _uri(shape) == URI_GRAFICO:
                    clase = "Gráfico"
                elif _uri(shape) == URI_OLE:
                    clase = "Objeto incrustado (Excel u otro)"
                elif shape.has_table:
                    textos.extend(_celdas(numero, shape, transf, en_grupo))
                    continue
                else:
                    avisos.append(f"«{shape.name}» (SmartArt u otro objeto) no se toma ni como texto ni como imagen.")
                    continue
            elif marcador in MARCADORES_IMAGEN:
                clase = "Marcador de imagen vacío"
            if clase:
                imagenes.append({
                    "id": f"img_{numero:02d}_{len(imagenes) + 1:02d}",
                    "shape_id": shape.shape_id,
                    "shape_nombre": shape.name,
                    "nombre": "",
                    "clase": clase + (" (en un grupo)" if en_grupo else ""),
                    "caja": transf.caja(shape.left, shape.top, shape.width, shape.height),
                })
                continue
            if not shape.has_text_frame:
                continue
            contenido = reemplazo.texto_cuadro(shape.text_frame)
            es_cuadro = shape._element.xpath("./p:nvSpPr/p:cNvSpPr/@txBox") == ["1"]
            # Formas decorativas sin texto (rectángulos, líneas) no se listan
            if not contenido.strip() and not es_cuadro and marcador is None:
                continue
            origen = (NOMBRES_MARCADOR.get(marcador, "Marcador") if marcador is not None
                      else "Cuadro de texto" if es_cuadro else "Forma con texto")
            textos.append(_texto(numero, shape, transf, en_grupo, shape.text_frame, origen))
        for hijo in slide.shapes._spTree.iter():
            if hijo.tag in DESCONOCIDOS:
                avisos.append(f"Hay un {DESCONOCIDOS[hijo.tag]} que la app no reconoce.")
        for i, img in enumerate(imagenes, 1):
            img["nombre"] = f"Diapositiva {numero} – imagen {i}"
        diapositivas.append({
            "numero": numero,
            "oculta": slide._element.get("show") == "0",
            "diseno": slide.slide_layout.name,
            "textos": textos,
            "imagenes": imagenes,
            "avisos": avisos,
        })
    return {
        "version": VERSION,
        "tamano_diapositiva": {"ancho": prs.slide_width, "alto": prs.slide_height,
                               "ancho_cm": cm(prs.slide_width), "alto_cm": cm(prs.slide_height)},
        "diapositivas": diapositivas,
    }


# --------------------------------------------------------------- buscar en el pptx
def buscar_shape(slide, shape_id: int = None, nombre: str = None):
    """Busca un shape de la diapositiva (también dentro de grupos) por id o por nombre."""
    for shape, _, _ in _recorrer(slide.shapes):
        if (nombre is not None and shape.name == nombre) or (shape_id is not None and shape.shape_id == shape_id):
            return shape
    return None


def text_frame_de(prs, numero: int, texto: dict):
    """El cuadro de texto (o la celda) de la plantilla dentro del PowerPoint."""
    shape = buscar_shape(prs.slides[numero - 1], shape_id=texto["shape_id"])
    if shape is None:
        raise ErrorPlantilla(f"No se encontró el texto {texto['id']} en la diapositiva {numero}.")
    if texto.get("celda"):
        r, c = texto["celda"]
        return shape.table.cell(r - 1, c - 1).text_frame
    return shape.text_frame


def lugar_de_imagen(prs, numero: int, imagen: dict):
    """El shape de un lugar de imagen: se busca por su identificador (nombre del shape)."""
    slide = prs.slides[numero - 1]
    return buscar_shape(slide, nombre=imagen["id"]) or buscar_shape(slide, shape_id=imagen["shape_id"])


# --------------------------------------------------------------------- variables
def variables(plantilla: dict) -> dict:
    """{id de variable: {nombre, usos: [(diapositiva, id de texto)]}} de toda la plantilla."""
    out = {}
    for d in plantilla["diapositivas"]:
        for t in d["textos"]:
            if t["tipo"] == VARIABLE and t["variable_id"]:
                v = out.setdefault(t["variable_id"], {"nombre": t["variable_nombre"], "usos": []})
                v["usos"].append((d["numero"], t["id"]))
    return out


def validar(plantilla: dict) -> list[str]:
    """Problemas que impiden guardar la plantilla (lista vacía = todo bien)."""
    errores = []
    if not plantilla.get("nombre", "").strip():
        errores.append("La plantilla no tiene nombre.")
    nombres_de = {}
    for d in plantilla["diapositivas"]:
        for t in d["textos"]:
            if t["tipo"] != VARIABLE:
                continue
            donde = f"Diapositiva {d['numero']}, texto «{corto(t['contenido'])}»"
            if not t["variable_nombre"].strip():
                errores.append(f"{donde}: falta el nombre de la variable.")
            if not re.fullmatch(r"[a-z0-9_]+", t["variable_id"] or ""):
                errores.append(f"{donde}: el identificador solo puede tener minúsculas, números y _.")
            if t["reemplazo"] == FRAGMENTO:
                if not t["fragmento"]:
                    errores.append(f"{donde}: falta escribir la parte a reemplazar.")
                elif t["fragmento"] not in t["contenido"]:
                    errores.append(f"{donde}: «{t['fragmento']}» no aparece en el texto (tiene que "
                                   "estar escrito igual, con mayúsculas y espacios).")
            if t["variable_id"]:
                nombres_de.setdefault(t["variable_id"], set()).add(t["variable_nombre"].strip())
    for vid, nombres in nombres_de.items():
        if len(nombres) > 1:
            errores.append(f"El identificador «{vid}» se usa con nombres distintos ({', '.join(sorted(nombres))}). "
                           "Si es el mismo dato, poné el mismo nombre; si no, cambiá el identificador.")
    for d in plantilla["diapositivas"]:
        for img in d["imagenes"]:
            if not img["nombre"].strip():
                errores.append(f"Diapositiva {d['numero']}: el lugar de imagen {img['id']} no tiene nombre.")
    return errores


def corto(texto: str, largo: int = 40) -> str:
    t = " ".join(texto.split())
    return t if len(t) <= largo else t[:largo - 1] + "…"


def rellenar_textos(prs, plantilla: dict, valores: dict) -> list[str]:
    """Pone en el PowerPoint el valor de cada variable (valores: {id de variable: texto}).

    Las variables sin valor quedan como estaban. Devuelve avisos (lo que no se pudo reemplazar).
    """
    avisos = []
    for d in plantilla["diapositivas"]:
        for t in d["textos"]:
            if t["tipo"] != VARIABLE or t["variable_id"] not in valores:
                continue
            tf = text_frame_de(prs, d["numero"], t)
            if t["reemplazo"] == FRAGMENTO:
                if reemplazo.reemplazar_fragmento(tf, t["fragmento"], valores[t["variable_id"]]) == 0:
                    avisos.append(f"Diapositiva {d['numero']}: no se encontró «{t['fragmento']}».")
            else:
                reemplazo.reemplazar_todo(tf, valores[t["variable_id"]])
    return avisos


def prueba_textos(plantilla: dict, raiz: Path = RAIZ) -> bytes:
    """PowerPoint de prueba: cada variable reemplazada por «su nombre», para revisar el formato."""
    prs = Presentation(str(carpeta(plantilla["id"], raiz) / "base.pptx"))
    rellenar_textos(prs, plantilla, {vid: f"«{v['nombre']}»" for vid, v in variables(plantilla).items()})
    salida = io.BytesIO()
    prs.save(salida)
    return salida.getvalue()


# -------------------------------------------------------------------- en disco
def carpeta(pid: str, raiz: Path = RAIZ) -> Path:
    return raiz / "plantillas" / pid


def carpeta_miniaturas(pid: str, raiz: Path = RAIZ) -> Path:
    return carpeta(pid, raiz) / "miniaturas"


def miniatura(pid: str, numero: int, raiz: Path = RAIZ) -> Path | None:
    ruta = carpeta_miniaturas(pid, raiz) / f"diapositiva_{numero:02d}.png"
    return ruta if ruta.exists() else None


def _nuevo_id(nombre: str, raiz: Path) -> str:
    while True:
        pid = f"{slug(nombre)[:40] or 'plantilla'}_{secrets.token_hex(3)}"
        if not carpeta(pid, raiz).exists():
            return pid


def _escribir_json(plantilla: dict, raiz: Path):
    (carpeta(plantilla["id"], raiz) / "plantilla.json").write_text(
        json.dumps(plantilla, ensure_ascii=False, indent=2), encoding="utf-8")


def _nombrar_lugares(prs, plantilla: dict):
    """Escribe el identificador de cada lugar de imagen como nombre del shape."""
    for d in plantilla["diapositivas"]:
        slide = prs.slides[d["numero"] - 1]
        for img in d["imagenes"]:
            shape = buscar_shape(slide, shape_id=img["shape_id"])
            if shape is not None:
                shape.name = img["id"]


def crear(nombre: str, archivo: str, contenido: bytes, raiz: Path = RAIZ) -> dict:
    """Crea una plantilla nueva a partir de un PowerPoint (todos los textos arrancan fijos)."""
    try:
        prs = Presentation(io.BytesIO(contenido))
    except Exception as e:
        raise ErrorPlantilla(f"No se pudo abrir «{archivo}» como PowerPoint (.pptx): {e}") from e
    if len(prs.slides) == 0:
        raise ErrorPlantilla("El PowerPoint no tiene diapositivas.")
    asegurar_ids_unicos(prs)
    plantilla = {"id": _nuevo_id(nombre, raiz), "nombre": nombre.strip() or Path(archivo).stem,
                 "archivo_original": archivo, "creada": ahora(), "modificada": ahora(), **analizar(prs)}
    carpeta(plantilla["id"], raiz).mkdir(parents=True)
    _nombrar_lugares(prs, plantilla)
    prs.save(str(carpeta(plantilla["id"], raiz) / "base.pptx"))
    _escribir_json(plantilla, raiz)
    return plantilla


def guardar(plantilla: dict, raiz: Path = RAIZ) -> dict:
    """Guarda los cambios (después de validar) y escribe los identificadores en base.pptx."""
    if errores := validar(plantilla):
        raise ErrorPlantilla("\n".join(errores))
    plantilla = {**plantilla, "nombre": plantilla["nombre"].strip(), "modificada": ahora()}
    base = carpeta(plantilla["id"], raiz) / "base.pptx"
    prs = Presentation(str(base))
    _nombrar_lugares(prs, plantilla)
    prs.save(str(base))
    _escribir_json(plantilla, raiz)
    return plantilla


def cargar(pid: str, raiz: Path = RAIZ) -> dict:
    return json.loads((carpeta(pid, raiz) / "plantilla.json").read_text(encoding="utf-8"))


def listar(raiz: Path = RAIZ) -> tuple[list[dict], list[str]]:
    """(plantillas guardadas, de la más nueva a la más vieja; carpetas que no se pudieron leer)."""
    plantillas, rotas = [], []
    base = raiz / "plantillas"
    for c in sorted(base.iterdir()) if base.exists() else []:
        if not c.is_dir():
            continue
        try:
            plantillas.append(cargar(c.name, raiz))
        except (OSError, ValueError, KeyError):
            rotas.append(c.name)
    plantillas.sort(key=lambda p: p.get("modificada", ""), reverse=True)
    return plantillas, rotas


def duplicar(pid: str, raiz: Path = RAIZ) -> dict:
    original = cargar(pid, raiz)
    nueva = {**original, "id": _nuevo_id(original["nombre"], raiz),
             "nombre": f"{original['nombre']} (copia)", "creada": ahora(), "modificada": ahora()}
    shutil.copytree(carpeta(pid, raiz), carpeta(nueva["id"], raiz))
    _escribir_json(nueva, raiz)
    return nueva


def eliminar(pid: str, raiz: Path = RAIZ):
    shutil.rmtree(carpeta(pid, raiz))


# ----------------------------------------------------------------------- resumen
def resumen(plantilla: dict) -> str:
    """Lo que se detectó en cada diapositiva, en texto (para revisar en la consola)."""
    tam = plantilla["tamano_diapositiva"]
    lineas = [f"Diapositivas: {len(plantilla['diapositivas'])}  ·  tamaño {tam['ancho_cm']} × {tam['alto_cm']} cm"]
    for d in plantilla["diapositivas"]:
        lineas.append(f"\n=== Diapositiva {d['numero']}{' (oculta)' if d['oculta'] else ''}"
                      f" · diseño «{d['diseno']}» · {len(d['textos'])} textos · {len(d['imagenes'])} imágenes")
        for i, t in enumerate(d["textos"], 1):
            c = t["caja"]
            lineas.append(f"  T{i} [{t['origen']}] ({cm(c['x'])}, {cm(c['y'])}) cm: "
                          f"«{corto(t['contenido'], 70) or '(vacío)'}»")
        for i, img in enumerate(d["imagenes"], 1):
            c = img["caja"]
            lineas.append(f"  I{i} [{img['clase']}] {img['id']} en ({cm(c['x'])}, {cm(c['y'])}) cm, "
                          f"{cm(c['ancho'])} × {cm(c['alto'])} cm  (antes se llamaba «{img['shape_nombre']}»)")
        lineas += [f"  ⚠ {a}" for a in d["avisos"]]
    return "\n".join(lineas)


if __name__ == "__main__":
    for ruta in sys.argv[1:]:
        p = Presentation(ruta)
        asegurar_ids_unicos(p)
        print(f"##### {ruta}\n{resumen(analizar(p))}\n")
