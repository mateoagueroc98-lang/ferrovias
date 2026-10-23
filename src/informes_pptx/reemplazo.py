"""Reemplazo de textos en un cuadro de PowerPoint conservando el formato.

PowerPoint parte el texto de un párrafo en varios "runs" (pedazos con su formato: fuente,
tamaño, color, negrita). Muchas veces los parte aunque el formato sea el mismo (por la
corrección ortográfica, por haber editado una letra...). Por eso "Septiembre 2026" puede estar
guardado como "Septi" + "embre 2026". Acá se busca sobre el texto del párrafo completo y el
reemplazo se escribe en el run donde empieza la coincidencia, con el formato de ese run.
"""
import copy
from dataclasses import dataclass

from pptx.oxml.ns import qn

_R, _BR, _FLD, _T = qn("a:r"), qn("a:br"), qn("a:fld"), qn("a:t")


@dataclass
class _Segmento:
    elem: object   # a:r, a:br o a:fld
    texto: str
    inicio: int    # posición dentro del texto del párrafo


def _segmentos(p) -> list[_Segmento]:
    """Pedazos de texto de un párrafo (a:p), en orden. Un salto de línea (a:br) cuenta como '\\n'."""
    segs, pos = [], 0
    for hijo in p:
        if hijo.tag in (_R, _FLD):
            t = hijo.find(_T)
            texto = (t.text or "") if t is not None else ""
        elif hijo.tag == _BR:
            texto = "\n"
        else:
            continue
        segs.append(_Segmento(hijo, texto, pos))
        pos += len(texto)
    return segs


def texto_parrafo(p) -> str:
    return "".join(s.texto for s in _segmentos(p))


def texto_cuadro(text_frame) -> str:
    """Texto completo del cuadro: párrafos separados por '\\n' (igual que lo ve el usuario)."""
    return "\n".join(texto_parrafo(p) for p in text_frame._txBody.findall(qn("a:p")))


def _poner_texto(elem, texto: str):
    t = elem.find(_T)
    if t is None:
        t = elem.makeelement(_T, {})
        elem.append(t)
    t.text = texto


def _reemplazar_rango(p, inicio: int, fin: int, nuevo: str):
    """Reemplaza los caracteres [inicio, fin) del párrafo por `nuevo`."""
    puesto = False
    for s in _segmentos(p):
        a, b = s.inicio, s.inicio + len(s.texto)
        if b <= inicio or a >= fin:
            continue
        if s.elem.tag == _BR:  # salto de línea dentro de lo reemplazado: se borra
            p.remove(s.elem)
            continue
        lo, hi = max(a, inicio) - a, min(b, fin) - a
        if not puesto:  # el run donde empieza la coincidencia recibe el texto nuevo (con su formato)
            texto = s.texto[:lo] + nuevo + s.texto[hi:]
            puesto = True
        else:           # los runs siguientes solo pierden la parte que ya se reemplazó
            texto = s.texto[:lo] + s.texto[hi:]
        if texto == "" and s.elem.tag == _R:
            p.remove(s.elem)
        else:
            _poner_texto(s.elem, texto)


def reemplazar_fragmento(text_frame, fragmento: str, nuevo: str) -> int:
    """Reemplaza cada aparición de `fragmento` (dentro de un mismo párrafo) por `nuevo`.

    Devuelve cuántas veces lo reemplazó. El texto nuevo toma el formato del run donde empezaba
    cada aparición.
    """
    if not fragmento:
        return 0
    total = 0
    for p in text_frame._txBody.findall(qn("a:p")):
        texto = texto_parrafo(p)
        inicios, i = [], texto.find(fragmento)
        while i >= 0:
            inicios.append(i)
            i = texto.find(fragmento, i + len(fragmento))
        # De atrás para adelante: así las posiciones de las anteriores no se corren
        for i in reversed(inicios):
            _reemplazar_rango(p, i, i + len(fragmento), nuevo)
        total += len(inicios)
    return total


def reemplazar_todo(text_frame, nuevo: str):
    """Reemplaza todo el texto del cuadro, con el formato del primer run del primer párrafo.

    Cada renglón de `nuevo` (separado por '\\n') queda en un párrafo con ese mismo formato.
    """
    cuerpo = text_frame._txBody
    parrafos = cuerpo.findall(qn("a:p"))
    if not parrafos:
        parrafos = [cuerpo.makeelement(qn("a:p"), {})]
        cuerpo.append(parrafos[0])
    p0 = parrafos[0]
    for p in parrafos[1:]:
        cuerpo.remove(p)

    run = p0.find(_R)
    fin_parrafo = p0.find(qn("a:endParaRPr"))
    if run is None:  # párrafo vacío: se crea un run con el formato de fin de párrafo
        run = p0.makeelement(_R, {})
        if fin_parrafo is not None:
            rpr = copy.deepcopy(fin_parrafo)
            rpr.tag = qn("a:rPr")
            run.append(rpr)
        run.append(run.makeelement(_T, {}))
        if fin_parrafo is not None:
            fin_parrafo.addprevious(run)
        else:
            p0.append(run)
    for hijo in list(p0):
        if hijo is not run and hijo.tag not in (qn("a:pPr"), qn("a:endParaRPr")):
            p0.remove(hijo)

    renglones = str(nuevo).split("\n")
    _poner_texto(run, renglones[0])
    anterior = p0
    for renglon in renglones[1:]:
        p = copy.deepcopy(p0)
        _poner_texto(p.find(_R), renglon)
        anterior.addnext(p)
        anterior = p
