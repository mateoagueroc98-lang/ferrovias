"""Concatenación: convierte el CSV del reporte LibroNovedades en un Excel ordenado.

El CSV trae una fila por cada RENGLÓN de texto (como el reporte paginado): número, estado,
fecha y hora se repiten en cada renglón. Acá se juntan en una fila por número de novedad,
con fecha, hora de inicio, hora de fin (renglón con palabra clave) y el detalle completo.
Todo local: sin IA ni APIs.
"""
import csv
import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ENCABEZADO = ["numero", "Descripción_Estado", "fecha", "hora", "descrip"]
COLUMNAS = ["Nro novedad", "Fecha", "Hora inicio", "Hora fin", "Detalle"]

ENCODINGS = {
    "Automático": None,  # UTF-8 si el archivo es UTF-8 válido; si no, Windows-1252
    "Windows-1252": "cp1252",
    "UTF-8": "utf-8-sig",
}

# Palabras clave que se ofrecen en pantalla. Solo NORMAL viene elegida.
CLAVES_DEFECTO = ["NORMAL"]
CLAVES_SUGERIDAS = ["NORMAL", "RECIBE VIA", "NORMALIZA", "NORMALIZADO", "NORMALIZADA",
                    "RESTABLECE", "RESTABLECIDO", "FUNCIONA CORRECTAMENTE", "SIN NOVEDAD"]

CONECTORES_DEFECTO = ["de", "del", "el", "la", "los", "las", "lo", "un", "una", "y", "e", "o",
                      "a", "al", "en", "que", "por", "con", "para", "se"]
# Cambian el sentido de la frase ("SEÑAL 23 NO ACLARA"): nunca son conectores.
NUNCA_CONECTORES = {"no", "sin"}

MAX_CELDA = 32767  # caracteres que admite una celda de Excel
FILAS_ENCABEZADO = 10  # el encabezado se busca entre las primeras filas


class ErrorArchivo(ValueError):
    """El archivo no tiene el formato esperado (mensaje listo para mostrar)."""


@dataclass
class Resultado:
    tabla: pd.DataFrame                    # columnas COLUMNAS, una fila por novedad
    renglon_fin: list[str]                 # por novedad, el renglón que marcó la hora fin ("" si ninguno)
    avisos: list[str] = field(default_factory=list)
    renglones: int = 0                     # renglones de datos leídos
    encoding: str = ""                     # encoding con el que se leyó


# ------------------------------------------------------------- normalización
def normalizar(texto: str) -> str:
    """'MEC. SILVA DA NORMAL.-' -> 'mec silva da normal' (minúsculas, sin tildes ni puntuación)."""
    t = unicodedata.normalize("NFD", str(texto).lower())
    t = "".join(ch for ch in t if not unicodedata.combining(ch))  # la ñ queda n
    return " ".join(re.sub(r"[\W_]+", " ", t).split())


def tokens(texto: str, conectores: set[str]) -> list[str]:
    """Palabras normalizadas del texto, sin conectores."""
    return [p for p in normalizar(texto).split() if p not in conectores]


def contiene(texto: list[str], clave: list[str]) -> bool:
    """¿La secuencia de palabras `clave` aparece junta y en orden dentro de `texto`?"""
    n = len(clave)
    return n > 0 and any(texto[i:i + n] == clave for i in range(len(texto) - n + 1))


def preparar_conectores(conectores: list[str]) -> tuple[set[str], list[str]]:
    """Conectores normalizados. Devuelve (conjunto, rechazados): "no" y "sin" se rechazan."""
    validos, rechazados = set(), []
    for c in conectores:
        n = normalizar(c)
        if not n:
            continue
        if n in NUNCA_CONECTORES or " " in n:  # un conector es una sola palabra
            rechazados.append(c)
        else:
            validos.add(n)
    return validos, rechazados


def preparar_claves(claves: list[str], conectores: set[str]) -> tuple[list[list[str]], list[str]]:
    """Palabras clave como listas de palabras. Devuelve (claves, ignoradas por quedar vacías)."""
    listas, ignoradas = [], []
    for c in claves:
        t = tokens(c, conectores)
        if t:
            if t not in listas:
                listas.append(t)
        else:
            ignoradas.append(c)
    return listas, ignoradas


# --------------------------------------------------------------------- lectura
def decodificar(contenido: bytes, encoding: str | None) -> tuple[str, str]:
    """Bytes -> (texto, nombre del encoding usado). encoding None = automático."""
    if encoding is None:
        try:
            return contenido.decode("utf-8-sig"), "UTF-8"
        except UnicodeDecodeError:
            encoding = "cp1252"
    try:
        texto = contenido.decode(encoding)
    except UnicodeDecodeError:
        if encoding == "cp1252":  # bytes que Windows-1252 no define: Latin-1 lee todo
            return contenido.decode("latin-1"), "Latin-1"
        raise ErrorArchivo("El archivo no se puede leer como UTF-8. "
                           "Probá con el encoding Windows-1252 o Automático.") from None
    return texto, "UTF-8" if encoding.startswith("utf") else "Windows-1252"


def _es_encabezado(fila: list[str]) -> bool:
    return [normalizar(c) for c in fila] == [normalizar(c) for c in ENCABEZADO]


def parsear_fecha(texto: str) -> dt.date | None:
    """'21/10/2026' -> date (siempre día primero). None si no es válida."""
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto.strip())
    if not m:
        return None
    try:
        return dt.date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def parsear_hora(texto: str) -> dt.time | None:
    """'03:16' (o '03:16:00') -> time. None si no es válida."""
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", texto.strip())
    if not m or int(m[1]) > 23 or int(m[2]) > 59:
        return None
    return dt.time(int(m[1]), int(m[2]))


def _lista(nums: list, max_items: int = 15) -> str:
    texto = ", ".join(str(n) for n in nums[:max_items])
    return texto + (f" y {len(nums) - max_items} más" if len(nums) > max_items else "")


def procesar(contenido: bytes, claves: list[str], conectores: list[str],
             encoding: str | None = None) -> Resultado:
    """Lee el CSV y arma la tabla con una fila por número de novedad.

    claves: palabras o frases que marcan el fin de la novedad (se buscan solo en `descrip`).
    conectores: palabras que se ignoran al comparar.
    encoding: "cp1252", "utf-8-sig" o None (automático).
    Lanza ErrorArchivo si el archivo no tiene el formato esperado.
    """
    texto, usado = decodificar(contenido, encoding)
    avisos = []
    conj_conectores, _ = preparar_conectores(conectores)
    listas_claves, _ = preparar_claves(claves, conj_conectores)

    filas = list(csv.reader(io.StringIO(texto)))
    inicio = None
    if any(any(c.strip() for c in f) for f in filas):
        inicio = next((i for i, f in enumerate(filas[:FILAS_ENCABEZADO]) if _es_encabezado(f)), None)
        if inicio is None:
            mensaje = ("No se encontró el encabezado esperado "
                       f"«{','.join(ENCABEZADO)}» en las primeras filas del archivo. "
                       "¿Es el CSV del reporte Libro de Novedades?")
            if "Ã" in texto:  # tildes rotas: UTF-8 leído como Windows-1252
                mensaje += " Parece que las tildes se leyeron mal: probá con el encoding UTF-8."
            raise ErrorArchivo(mensaje)

    # numero -> renglones del grupo (fila del CSV, fecha, hora, descrip), en orden de aparición
    grupos: dict[int, list[tuple[int, str, str, str]]] = {}
    salteadas, separados, ultimo = [], [], None
    for i, fila in enumerate(filas[inicio + 1:] if inicio is not None else [], start=(inicio or 0) + 2):
        if not any(c.strip() for c in fila):
            continue
        if len(fila) < 5:
            salteadas.append(i)
            continue
        if len(fila) > 5:  # coma sin comillas dentro del texto: se junta de nuevo
            fila = fila[:4] + [",".join(fila[4:])]
        try:
            numero = int(fila[0].strip())
        except ValueError:
            salteadas.append(i)
            continue
        if numero in grupos and numero != ultimo and numero not in separados:
            separados.append(numero)
        grupos.setdefault(numero, []).append((i, fila[2], fila[3], fila[4]))
        ultimo = numero

    renglones = sum(len(g) for g in grupos.values())
    filas_tabla, renglon_fin = [], []
    fechas_mal, horas_mal, truncadas = [], [], []
    for numero, grupo in grupos.items():
        i, f, _, _ = grupo[0]  # la fecha sale de la primera fila del grupo
        if parsear_fecha(f) is None:
            fechas_mal.append(f"fila {i} (nro {numero}: «{f}»)")
        for i, _, h, _ in grupo:
            if parsear_hora(h) is None:
                horas_mal.append(f"fila {i} (nro {numero}: «{h}»)")
        fin, texto_fin = None, ""
        for _, _, h, d in grupo:
            if any(contiene(tokens(d, conj_conectores), c) for c in listas_claves):
                fin, texto_fin = parsear_hora(h), d.strip()  # si coinciden varias, queda la última
        detalle = "\n".join(d.strip() for _, _, _, d in grupo if d.strip())
        if len(detalle) > MAX_CELDA:
            detalle = detalle[:MAX_CELDA]
            truncadas.append(numero)
        filas_tabla.append({
            "Nro novedad": numero,
            "Fecha": parsear_fecha(grupo[0][1]),
            "Hora inicio": parsear_hora(grupo[0][2]),
            "Hora fin": fin,
            "Detalle": detalle,
        })
        renglon_fin.append(texto_fin)

    if not grupos:
        avisos.append("El archivo no tiene novedades: el Excel sale solo con los encabezados.")
    if salteadas:
        avisos.append(f"Se saltearon {len(salteadas)} filas sin número de novedad válido o con "
                      f"columnas de menos: filas {_lista(salteadas)}.")
    if separados:
        avisos.append("Estos números de novedad aparecen separados por otros en el CSV (se juntaron "
                      f"igual en una sola fila): {_lista(separados)}.")
    if fechas_mal:
        avisos.append(f"Fecha con formato inválido (la celda queda vacía): {_lista(fechas_mal)}.")
    if horas_mal:
        avisos.append(f"Hora con formato inválido (la celda queda vacía): {_lista(horas_mal)}.")
    if truncadas:
        avisos.append(f"Detalle demasiado largo para una celda de Excel ({MAX_CELDA} caracteres), "
                      f"se recortó en las novedades: {_lista(truncadas)}.")

    tabla = pd.DataFrame(filas_tabla, columns=COLUMNAS)
    return Resultado(tabla, renglon_fin, avisos, renglones, usado)


# ------------------------------------------------------------------ exportación
def exportar_excel(tabla: pd.DataFrame) -> bytes:
    """Una hoja con encabezados en negrita, filtros, fecha/hora reales y Detalle con ajustar texto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Libro de Novedades"
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    formatos = {"Fecha": "dd/mm/yyyy", "Hora inicio": "hh:mm", "Hora fin": "hh:mm"}
    arriba = Alignment(vertical="top")
    ajustar = Alignment(vertical="top", wrap_text=True)
    for fila in tabla.itertuples(index=False):
        valores = [None if pd.isna(v) else v for v in fila]
        valores[0] = int(valores[0])
        ws.append(valores)
        for nombre, celda in zip(COLUMNAS, ws[ws.max_row]):
            if nombre in formatos:
                celda.number_format = formatos[nombre]
            celda.alignment = ajustar if nombre == "Detalle" else arriba
    for i, ancho in enumerate([12, 12, 11, 11, 100], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
