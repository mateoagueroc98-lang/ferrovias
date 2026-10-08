"""Genera dos Excel de prueba INVENTADOS para el módulo Asociación LG ↔ PCZ.

Uso:  python tests/generar_ejemplos.py
Crea tests/ejemplos/LG_ejemplo.xlsx y tests/ejemplos/PCZ_ejemplo.xlsx.
Las fechas vienen en formatos mezclados a propósito (fecha de Excel, número serial,
texto "DD/MM/AAAA" y "D/M/AA") y hay una fecha inválida para ver la advertencia.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook

CARPETA = Path(__file__).parent / "ejemplos"

# 04/03/2025, 05/03/2025 y 06/03/2025 (si se leyeran mes primero serían abril/mayo/junio)
D4, D5, D6 = dt.datetime(2025, 3, 4), dt.datetime(2025, 3, 5), dt.datetime(2025, 3, 6)
SERIAL_D5 = (dt.date(2025, 3, 5) - dt.date(1899, 12, 30)).days  # 45721

LG = [
    # Página, Fecha, Hora, Transcripción (un aviso por renglón)
    (12, D4, dt.time(8, 15),
     "Se informa falla en señal 14 de estación Boulogne, concurre mecánico Pérez.\n"
     "Se normaliza señal 14 a las 9:10 hs."),
    (12, "04/03/2025", "14:40",
     "Brazo del P.A.N. calle Sarmiento no baja, interviene Gómez.\n"
     "Cambio 7 sin comprobación en Villa Adelina, se ajusta."),
    (13, "5/3/25", dt.time(23, 50),
     "Circuito de vía 22 ocupado en km 14/116, se revisa junta aislada.\n"
     "Tren 3215 detenido por señal 31 a peligro."),
    (13, SERIAL_D5, "10:05",
     "Locomotora E712 informa barrera baja sin tren en P.A.N. Av. Márquez."),
    (14, D6, dt.time(6, 30),
     "Abrigo 3 sin energía en estación Don Torcuato, se cambia fusible.\n"
     "Se normaliza a las 7:05."),
]

PCZ = [
    # Descripción, Fecha
    ("Falla señal 14 Boulogne, a peligro 08:10 hs", D4),
    ("Señal 14 Boulogne normalizada 09:12", "04/03/2025"),
    ("Brazo P.A.N. Sarmiento no baja 14:35", "4/3/25"),
    ("Cambio 7 Villa Adelina sin comprobación 14:50", D4),
    ("Corte de energía en estación Retiro 03:00", D4),
    ("Tren 3101 demorado por pasajero indispuesto 07:20", "04/03/2025"),
    ("Vandalismo en estación Grand Bourg", D4),
    ("Paso a nivel Av. Márquez barrera rota 16:00", D4),
    ("Ocupación indebida de vía km 20/300", "04/03/2025"),
    ("Tren 3150 cancelado por falta de material rodante", D4),
    ("Señal 5 Villa Rosa a peligro 12:00", "04/03/2025"),
    ("Circuito 22 ocupado km 14/116 23:45", SERIAL_D5),
    ("Tren 3215 detenido señal 31 peligro 23:55", "05/03/2025"),
    ("Barrera baja sin tren P.A.N. Márquez, informa loc E-712 10:00", "5/3/25"),
    ("Cambio 12 Boulogne trabado 11:30", D5),
    ("Falla de comunicación radio tren 3220", D5),
    ("Robo de cables en km 25/100", "05/03/2025"),
    ("Pasajero accidentado estación Tortuguitas 18:20", D5),
    ("Señal 9 Pilar sin luz 20:10", SERIAL_D5),
    ("Corte de alimentación en PAN calle Belgrano 13:00", "05/03/2025"),
    ("Tren 3230 demorado 15 minutos por congestión", D5),
    ("Abrigo 3 Don Torcuato sin energía 06:20", D6),
    ("Normalizado abrigo 3 Don Torcuato 07:00", "06/03/2025"),
    ("Circuito de vía 22 normalizado 00:40", "6/3/25"),
    ("Señal 14 Boulogne intermitente 10:30", D6),
    ("Tren 3301 cancelado", "06/03/2025"),
    ("Paso a nivel calle Sarmiento barreras altas 12:15", D6),
    ("Animal sobre vía km 30/200", D6),
    ("Cambio 7 Villa Adelina lubricado", "06/03/2025"),
    ("Novedad con fecha mal cargada", "32/13/2025"),
]


def generar(carpeta: Path = CARPETA) -> tuple[Path, Path]:
    carpeta.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.append(["Pagina", "Fecha", "Hora", "Transcripcion"])
    for fila in LG:
        ws.append(list(fila))
    for celda in ws["B"][1:]:
        if isinstance(celda.value, dt.datetime):
            celda.number_format = "DD/MM/YYYY"
    ruta_lg = carpeta / "LG_ejemplo.xlsx"
    wb.save(ruta_lg)

    wb = Workbook()
    ws = wb.active
    ws.append(["Descripcion", "Fecha", "Nº Novedad"])
    for n, (desc, fecha) in enumerate(PCZ, start=1001):
        ws.append([desc, fecha, n])
    for celda in ws["B"][1:]:
        if isinstance(celda.value, dt.datetime):
            celda.number_format = "DD/MM/YYYY"
    ruta_pcz = carpeta / "PCZ_ejemplo.xlsx"
    wb.save(ruta_pcz)
    return ruta_lg, ruta_pcz


if __name__ == "__main__":
    for r in generar():
        print("Creado:", r)
