"""Tests del módulo Asociación LG ↔ PCZ (correr con:  python -m pytest)."""
import datetime as dt
import io

import pandas as pd
import pytest

from src import asociacion_lg_pcz as a
from tests.generar_ejemplos import generar

F = dt.date


# ------------------------------------------------------------- fechas
@pytest.mark.parametrize("valor, esperado", [
    ("05/03/2025", F(2025, 3, 5)),      # día primero, nunca mes primero
    ("5/3/25", F(2025, 3, 5)),
    ("05-03-2025", F(2025, 3, 5)),
    ("5.3.2025", F(2025, 3, 5)),
    ("13/01/2025", F(2025, 1, 13)),
    (" 05/03/2025 14:30 ", F(2025, 3, 5)),
    ("2025-03-05", F(2025, 3, 5)),      # ISO
    ("2025-03-05 00:00:00", F(2025, 3, 5)),
    (45721, F(2025, 3, 5)),             # número serial de Excel
    (45721.75, F(2025, 3, 5)),          # serial con hora
    ("45721", F(2025, 3, 5)),
    (dt.datetime(2025, 3, 5, 10, 0), F(2025, 3, 5)),
    (pd.Timestamp("2025-03-05"), F(2025, 3, 5)),
    (F(2025, 3, 5), F(2025, 3, 5)),
])
def test_parsear_fecha_valida(valor, esperado):
    assert a.parsear_fecha(valor) == esperado


@pytest.mark.parametrize("valor", [
    "03/13/2025",   # mes 13: NO se reinterpreta como MM/DD
    "32/01/2025", "31/02/2025", "ayer", "", "  ", None, float("nan"), pd.NaT, 12, "1/2",
])
def test_parsear_fecha_invalida(valor):
    assert a.parsear_fecha(valor) is None


@pytest.mark.parametrize("valor, esperado", [
    ("14:40", 14 * 60 + 40), ("9.05", 9 * 60 + 5), ("22hs", 22 * 60), ("08:15:00", 8 * 60 + 15),
    (dt.time(23, 50), 23 * 60 + 50), (0.5, 12 * 60), (None, None), ("sin hora", None),
])
def test_parsear_hora(valor, esperado):
    assert a.parsear_hora(valor) == esperado


def test_horas_en_texto_no_confunde_km():
    horas = [h for _, _, h in a.horas_en_texto("km 14/116 a las 9:10 hs, km 14.300, 22 hs")]
    assert horas == [9 * 60 + 10, 22 * 60]


# ------------------------------------------------------------ columnas
def test_mapear_columnas_tolerante():
    df = pd.DataFrame(columns=["  DESCRIPCION ", "fecha", "Nº Novedad", "Hora del evento"])
    mapa, faltan = a.mapear_columnas(df, a.COLUMNAS_PCZ)
    assert faltan == list(a.EXTRAS_PCZ) and mapa["Hora"] == "Hora del evento"
    assert mapa["N° Novedad"] == "Nº Novedad"
    for nombre in ["N° Novedad", "N Novedad", "Nro. Novedad", "nº novedad"]:
        _, faltan = a.mapear_columnas(pd.DataFrame(columns=["Descripción", "Fecha", nombre]), a.COLUMNAS_PCZ)
        assert faltan == ["Hora", *a.EXTRAS_PCZ], nombre  # hora y columnas extra son opcionales


def test_mapear_columnas_avisa_faltantes():
    _, faltan = a.mapear_columnas(pd.DataFrame(columns=["Pagina", "Fecha"]), a.COLUMNAS_LG)
    assert faltan == ["Transcripción", "Hora"]


# ------------------------------------------------------ filtro por fecha
def _pcz():
    df = pd.DataFrame({
        "Descripción": ["a", "b", "c", "d", "e"],
        "Fecha": ["04/03/2025", 45721, "6/3/25", "05/03/2025", "basura"],
        "N° Novedad": [1, 2, 3, 4, 5],
    })
    mapa, _ = a.mapear_columnas(df, a.COLUMNAS_PCZ)
    return a.preparar_pcz(df, mapa)


def test_filtro_solo_misma_fecha():
    sel = a.filtrar_por_fecha(_pcz(), F(2025, 3, 5))
    assert list(sel["id"]) == ["2", "4"]
    assert not sel["otro_dia"].any()


def test_filtro_no_confunde_mes_y_dia():
    # 3 de mayo NO es 5 de marzo
    assert a.filtrar_por_fecha(_pcz(), F(2025, 5, 3)).empty


def test_filtro_con_dias_vecinos_marcados():
    sel = a.filtrar_por_fecha(_pcz(), F(2025, 3, 5), vecinos=True)
    assert sorted(sel["id"]) == ["1", "2", "3", "4"]
    assert dict(zip(sel["id"], sel["otro_dia"])) == {"1": True, "2": False, "3": True, "4": False}


def test_filtro_sin_fecha_lg_no_da_candidatos():
    assert a.filtrar_por_fecha(_pcz(), None).empty


def test_fechas_invalidas_se_informan():
    malas = a.fechas_invalidas(_pcz())
    assert list(malas["Fila"]) == [5]
    assert list(malas["Fecha tal cual viene"]) == ["basura"]


# ------------------------------------------------------------- puntaje
def test_similitud_numeros_palabras_y_hora():
    r = a.similitud("Falla en señal 14 de estación Boulogne", [8 * 60 + 15],
                    "Señales 14 Boulogne a peligro 08:10 hs")
    assert "n:14" in r["comunes"] and "w:boulo" in r["comunes"] and "w:senal" in r["comunes"]
    assert r["dif_hora"] == 5
    assert r["puntaje"] == 3 + 2 + 3


def test_similitud_sigla_y_km():
    r = a.similitud("Locomotora E712 en km 14/116", [], "loc E-712 km 14/116")
    assert {"n:712", "n:14/116"} <= r["comunes"]


def test_hora_cruza_medianoche():
    r = a.similitud("x", [23 * 60 + 50], "y 00:10")
    assert r["dif_hora"] == 20


def test_resaltar_escapa_html():
    salida = a.resaltar("Señal 14 <b>", {"n:14"})
    assert "<mark>14</mark>" in salida and "&lt;b&gt;" in salida


def test_candidatos_ordenados_y_todos_visibles():
    lg = pd.DataFrame({"fecha": [F(2025, 3, 5), F(2025, 3, 5), F(2025, 3, 4)], "hora": [None] * 3,
                       "avisos": [["Tren 3101 demorado"], ["Falla señal 14", "Cambio 7"], ["Señal 14"]]})
    avisos = a.tabla_avisos(lg)
    assert list(avisos["clave"]) == ["0:0", "1:0", "1:1", "2:0"]
    nov = pd.Series({"fecha": F(2025, 3, 5), "hora": None, "descripcion": "Señal 14 a peligro"})
    c = a.candidatos(nov, avisos)
    assert list(c["clave"]) == ["1:0", "0:0", "1:1"]  # todos los avisos del día, el mejor primero


# ---------------------------------------- deshacer, progreso y exportación
def _cargar_ejemplos(tmp_path):
    ruta_lg, ruta_pcz = generar(tmp_path)
    lg_df = a.leer_tabla(ruta_lg.name, ruta_lg.read_bytes())
    pcz_df = a.leer_tabla(ruta_pcz.name, ruta_pcz.read_bytes())
    lg = a.preparar_lg(lg_df, a.mapear_columnas(lg_df, a.COLUMNAS_LG)[0])
    pcz = a.preparar_pcz(pcz_df, a.mapear_columnas(pcz_df, a.COLUMNAS_PCZ)[0])
    return lg_df, pcz_df, lg, pcz


def test_ejemplos_fechas_mezcladas(tmp_path):
    _, _, lg, pcz = _cargar_ejemplos(tmp_path)
    assert lg["fecha"].notna().all()
    assert set(lg["fecha"]) == {F(2025, 3, 4), F(2025, 3, 5), F(2025, 3, 6)}
    assert pcz["fecha"].isna().sum() == 1  # la "32/13/2025"
    # Para la novedad 1014 (E-712, Márquez), el mejor aviso es el de la fila 4 de LG
    nov = pcz[pcz["id"] == "1014"].iloc[0]
    assert a.candidatos(nov, a.tabla_avisos(lg)).iloc[0]["clave"] == "3:0"


def test_dias_estado_y_progreso():
    pcz = pd.DataFrame({"fecha": [F(2025, 3, 5), F(2025, 3, 4), F(2025, 3, 5), None]})
    assert a.dias_pcz(pcz) == [F(2025, 3, 4), F(2025, 3, 5), None]
    assert a.novedades_del_dia(pcz, F(2025, 3, 5)) == [0, 2] and a.novedades_del_dia(pcz, None) == [3]
    s = a.trabajo_nuevo("h")
    a.asociar(s, [0, 2], ["0:0", "1:1"])
    a.descartar(s, [1])
    assert [a.estado(s, k) for k in range(4)] == ["asociada", "descartada", "asociada", "pendiente"]
    assert a.progreso(s, 4) == {"total": 4, "asociadas": 2, "descartadas": 1, "pendientes": 1}
    assert a.pendientes(s, [0, 1, 2, 3]) == [3]
    assert a.novedades_por_aviso(s) == {"0:0": [0, 2], "1:1": [0, 2]}
    a.recuperar(s, 1); a.quitar_asociacion(s, 0)
    assert a.pendientes(s, [0, 1, 2, 3]) == [0, 1, 2, 3]


def test_novedades_unidas_puntuan_juntas():
    pcz = pd.DataFrame({"fecha": [F(2025, 3, 4)] * 2, "descripcion": ["Señal 14", "Boulogne"],
                        "hora": [None, 8 * 60 + 10]})
    nov = a.novedad_combinada(pcz, [0, 1])
    assert nov["descripcion"] == "Señal 14 Boulogne" and nov["hora"] == [8 * 60 + 10]
    lg = pd.DataFrame({"fecha": [F(2025, 3, 4)] * 2, "hora": [8 * 60 + 15, None],
                       "avisos": [["Falla señal 14 Boulogne"], ["Cambio 7"]]})
    c = a.candidatos(nov, a.tabla_avisos(lg))
    assert list(c["clave"]) == ["0:0", "1:0"] and "hora a 5 min" in c.iloc[0]["detalle"]


def _trabajo_ejemplo(pcz):
    s = a.trabajo_nuevo("h")
    pos = {n: k for k, n in enumerate(pcz["id"])}
    # 1001 y 1002 unidas con los dos avisos de la fila 1 de LG
    a.asociar(s, [pos["1001"], pos["1002"]], ["0:0", "0:1"])
    a.asociar(s, [pos["1003"]], ["1:0"])          # 1003: aviso 1 de la fila 2
    a.asociar(s, [pos["1022"]], ["4:0", "0:0"])   # 1022: avisos de las filas 5 y 1
    a.descartar(s, [pos["1004"]])
    return s


MAPA_LG = {"Fecha": "Fecha", "Hora": "Hora", "Transcripción": "Transcripcion"}
MAPA_PCZ = {"Fecha": "Fecha", "Descripción": "Descripcion", "N° Novedad": "Nº Novedad", "Hora": "Hora"}


def test_estructura_fija_de_la_hoja_asociaciones():
    assert list(a.COLUMNAS_SALIDA) == [
        "N°", "ID", "Fecha", "Mes", "N° Novedad", "Sistema", "Subsistema", "Parte", "Elemento",
        "Nombre correcto", "Km. INPAN", "BuscarV (Km. INPAN)", "Modelo", "Descripción",
        "Zona", "VER", "Falla", "Comienzo de falla", "Habilitación"]
    assert a.texto_descripcion(["Falla señal 14", "Normalizada"], ["Se informa falla.", "Se normaliza."]) == (
        "FALLA SEÑAL 14\nNORMALIZADA\n\nSe informa falla.\nSe normaliza.")
    assert a.texto_descripcion(["Falla señal 14"], []) == "FALLA SEÑAL 14"  # sin asociar: solo PCZ


def test_tabla_asociaciones_trae_todas_las_novedades_por_dia(tmp_path):
    _, _, lg, pcz = _cargar_ejemplos(tmp_path)
    s = _trabajo_ejemplo(pcz)
    t = a.tabla_asociaciones(lg, pcz, s, numero_inicial=120)
    # 30 novedades: 4 en 3 asociaciones (1001+1002 unidas) y 26 sin asociar, una fila cada una
    assert list(t.columns) == list(a.COLUMNAS_SALIDA) and len(t) == 29
    assert list(t["N°"]) == list(range(120, 149))  # consecutivo desde el número elegido
    ids = [str(x) for x in t["N° Novedad"]]
    assert ids[:3] == ["1001 | 1002", "1003", "1004"] and ids[19:21] == ["1021", "1022"]
    assert ids[-1] == "1030"  # sin fecha reconocida: al final
    assert t.iloc[2]["Descripción"] == "CAMBIO 7 VILLA ADELINA SIN COMPROBACIÓN 14:50"  # descartada: solo PCZ
    vacias = [c for c, uso in a.COLUMNAS_SALIDA.items() if uso is None]
    assert t[vacias].isna().all().all()
    f = t.iloc[0]  # 1001 + 1002 unidas
    assert f["N° Novedad"] == "1001 | 1002" and f["Fecha"] == F(2025, 3, 4)
    assert f["Descripción"] == ("FALLA SEÑAL 14 BOULOGNE, A PELIGRO 08:10 HS\nSEÑAL 14 BOULOGNE NORMALIZADA 09:12\n\n"
                                "Se informa falla en señal 14 de estación Boulogne, concurre mecánico Pérez.\n"
                                "Se normaliza señal 14 a las 9:10 hs.")
    assert t.iloc[1]["N° Novedad"] == 1003  # una sola novedad: número


def test_columnas_de_pcz_despues_de_la_descripcion(tmp_path):
    _, _, lg, pcz = _cargar_ejemplos(tmp_path)
    t = a.tabla_asociaciones(lg, pcz, _trabajo_ejemplo(pcz))
    unidas, sola = t.iloc[0], t.iloc[2]  # 1001 + 1002 unidas; 1004 sola (descartada)
    assert unidas["Zona"] == "Zona 3 | Zona 1" and unidas["Falla"] == "Barreras | Energía"
    assert sola["Zona"] == "Zona 3" and sola["VER"] == "No" and sola["Falla"] == "Señalamiento"
    assert sola["Comienzo de falla"] == pd.Timestamp(2025, 3, 4, 14, 50)  # fecha y hora, sin tocar
    assert sola["Habilitación"] == pd.Timestamp(2025, 3, 4, 15, 35)


def test_columnas_de_pcz_vacias_si_el_archivo_no_las_tiene():
    df = pd.DataFrame({"Fecha": ["04/03/2025"], "Descripcion": ["A"]})
    pcz = a.preparar_pcz(df, {"Fecha": "Fecha", "Descripción": "Descripcion"})
    lg = a.preparar_lg(pd.DataFrame({"Fecha": ["04/03/2025"], "T": ["X"]}), {"Fecha": "Fecha", "Transcripción": "T"})
    t = a.tabla_asociaciones(lg, pcz, a.trabajo_nuevo("h"))
    assert t[list(a.EXTRAS_PCZ)].isna().all().all()


def test_tabla_asociaciones_ordena_por_dia_aunque_el_pcz_venga_desordenado():
    df = pd.DataFrame({"Fecha": ["06/03/2025", "04/03/2025", "sin fecha", "04/03/2025", "05/03/2025"],
                       "Descripcion": ["A", "B", "C", "D", "E"], "N": [1, 2, 3, 4, 5]})
    pcz = a.preparar_pcz(df, {"Fecha": "Fecha", "Descripción": "Descripcion", "N° Novedad": "N"})
    lg = a.preparar_lg(pd.DataFrame({"Fecha": ["04/03/2025"], "T": ["Aviso"]}), {"Fecha": "Fecha", "Transcripción": "T"})
    s = a.trabajo_nuevo("h")
    a.asociar(s, [3], ["0:0"])                    # D (4/3) asociada
    t = a.tabla_asociaciones(lg, pcz, s)
    assert list(t["N° Novedad"]) == [2, 4, 5, 1, 3]  # 4/3 (en orden del PCZ), 5/3, 6/3, sin fecha
    assert t.iloc[1]["Descripción"] == "D\n\nAviso"


def test_sin_asociar_conserva_columnas_originales(tmp_path):
    lg_df, pcz_df, lg, pcz = _cargar_ejemplos(tmp_path)
    s = _trabajo_ejemplo(pcz)
    lgs = a.lg_sin_asociar(lg_df, lg, s, MAPA_LG)
    assert list(lgs.columns) == list(lg_df.columns)
    # fila 1 completa; fila 2 solo su aviso 2; filas 3 y 4 enteras; fila 5 solo su aviso 2
    assert list(lgs["Pagina"]) == [12, 13, 13, 14]
    assert lgs.iloc[0]["Transcripcion"] == "Cambio 7 sin comprobación en Villa Adelina, se ajusta."
    assert lgs.iloc[3]["Transcripcion"] == "Se normaliza a las 7:05."


def test_exportar_dos_hojas(tmp_path):
    lg_df, pcz_df, lg, pcz = _cargar_ejemplos(tmp_path)
    s = _trabajo_ejemplo(pcz)
    datos = a.exportar(lg_df, lg, pcz, s, MAPA_LG, numero_inicial=7)
    hojas = pd.read_excel(io.BytesIO(datos), sheet_name=None)
    assert list(hojas) == ["Asociaciones", "LG sin asociar"]
    assert list(hojas["Asociaciones"].columns) == list(a.COLUMNAS_SALIDA) and len(hojas["Asociaciones"]) == 29
    assert len(hojas["LG sin asociar"]) == 4
    # En el Excel: fecha real con formato d/m/aaaa y la descripción con «ajustar texto»
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(datos))["Asociaciones"]
    assert hoja["A2"].value == 7 and hoja["C2"].number_format == "d/m/yyyy"
    assert hoja["C2"].value.date() == F(2025, 3, 4)
    assert hoja["N2"].alignment.wrap_text and hoja["N2"].value.startswith("FALLA SEÑAL 14")
    assert hoja["O4"].value == "Zona 3" and hoja["R4"].number_format == "d/m/yyyy h:mm"  # 1004


def test_marcar_pone_cambia_y_saca():
    s = a.trabajo_nuevo("h")
    a.marcar(s, 3, "amarilla")
    assert s["marcas"] == {3: "amarilla"}
    a.marcar(s, 3, "naranja")                     # otro color: lo reemplaza
    assert s["marcas"] == {3: "naranja"}
    a.marcar(s, 3, "naranja")                     # el mismo color otra vez: la saca
    assert s["marcas"] == {}


def test_marcar_una_marca_tambien_las_unidas():
    s = a.trabajo_nuevo("h")
    a.marcar(s, 1, "amarilla", juntas=[1, 2, 3])
    assert s["marcas"] == {1: "amarilla", 2: "amarilla", 3: "amarilla"}
    a.marcar(s, 2, "naranja", juntas=[1, 2, 3])   # cambiar el color de una cambia todas
    assert s["marcas"] == {1: "naranja", 2: "naranja", 3: "naranja"}
    a.marcar(s, 3, "naranja", juntas=[1, 2, 3])   # sacarla a una la saca de todas
    assert s["marcas"] == {}


def test_asociar_unifica_las_marcas():
    s = a.trabajo_nuevo("h")
    a.marcar(s, 1, "amarilla")
    a.marcar(s, 2, "naranja")
    a.asociar(s, [1, 2, 3], ["0:0"])
    assert s["marcas"] == {1: "naranja", 2: "naranja", 3: "naranja"}


def test_marca_de_una_fila_naranja_pesa_mas():
    s = a.trabajo_nuevo("h")
    a.marcar(s, 1, "amarilla")
    a.marcar(s, 2, "naranja")
    assert a.marca(s, [1]) == "amarilla" and a.marca(s, [1, 2]) == "naranja" and a.marca(s, [5]) is None


def test_exportar_resalta_las_filas_marcadas(tmp_path):
    from openpyxl import load_workbook
    lg_df, pcz_df, lg, pcz = _cargar_ejemplos(tmp_path)
    s = _trabajo_ejemplo(pcz)
    pos = {n: k for k, n in enumerate(pcz["id"])}
    a.marcar(s, pos["1002"], "amarilla")          # asociada (unida con 1001): fila 2 de Asociaciones
    a.marcar(s, pos["1003"], "naranja")           # asociada: fila 3 de Asociaciones
    a.marcar(s, pos["1004"], "amarilla")          # descartada: fila 4 (también está en la hoja)
    hoja = load_workbook(io.BytesIO(a.exportar(lg_df, lg, pcz, s, MAPA_LG)))["Asociaciones"]
    color = lambda celda: celda.fill.fgColor.rgb if celda.fill.fill_type == "solid" else None
    assert {color(c) for c in hoja[2]} == {"00FFFF00"}
    assert {color(c) for c in hoja[3]} == {"00FFC000"}
    assert {color(c) for c in hoja[4]} == {"00FFFF00"}
    assert all(color(c) is None for fila in hoja.iter_rows(min_row=5) for c in fila)


# ------------------------------------------------ grupos aviso(s) -> novedad(es)
def test_puntaje_por_aviso_y_no_por_parrafo():
    lg = pd.DataFrame({"fecha": [F(2025, 3, 4)], "hora": [None],
                       "avisos": [["Señal 14 Boulogne", "Cambio 7 Villa Adelina"]]})
    nov = pd.Series({"fecha": F(2025, 3, 4), "hora": None, "descripcion": "Cambio 7 Villa Adelina"})
    c = a.candidatos(nov, a.tabla_avisos(lg))
    assert list(c["clave"]) == ["0:1", "0:0"] and c.iloc[1]["puntaje"] == 0


# ---------------------------------------------- columnas elegidas por el usuario
def test_preparar_lg_con_columnas_de_otro_nombre_y_sin_opcionales():
    df = pd.DataFrame({"Día": ["05/03/2025"], "Texto del libro": ["Señal 14\nCambio 7"]})
    lg = a.preparar_lg(df, {"Fecha": "Día", "Transcripción": "Texto del libro"})
    assert lg.loc[0, "fecha"] == F(2025, 3, 5)
    assert lg.loc[0, "avisos"] == ["Señal 14", "Cambio 7"]
    assert lg.loc[0, "pagina"] == "" and lg.loc[0, "hora"] is None


def test_preparar_pcz_con_columnas_de_otro_nombre():
    df = pd.DataFrame({"Detalle": ["Falla"], "F. evento": [45721], "ID": [7.0]})
    pcz = a.preparar_pcz(df, {"Descripción": "Detalle", "Fecha": "F. evento", "N° Novedad": "ID"})
    assert list(pcz["id"]) == ["7"] and pcz.loc[0, "fecha"] == F(2025, 3, 5)


def test_buscar_en_avisos():
    lg = pd.DataFrame({"fecha": [F(2025, 3, 4)], "hora": [None], "avisos": [["Señal 14 Boulogne", "Cambio 7"]]})
    av = a.tabla_avisos(lg)
    assert list(a.buscar(av, "senal boulogne")["clave"]) == ["0:0"]


def test_preparar_pcz_sin_n_novedad_usa_fila_del_excel():
    df = pd.DataFrame({"Detalle": ["a", "b"], "Día": ["04/03/2025", "05/03/2025"]})
    pcz = a.preparar_pcz(df, {"Descripción": "Detalle", "Fecha": "Día"})
    assert list(pcz["id"]) == ["2", "3"]


def test_hora_con_filas_vacias_se_muestra_bien():
    # Si algunas filas no tienen hora, pandas pasaba las horas a decimales (510.0) y fallaba
    df = pd.DataFrame({"Fecha": ["04/03/2025", "04/03/2025"], "Hora": ["08:30", None],
                       "Transcripción": ["a", "b"]})
    lg = a.preparar_lg(df, {"Fecha": "Fecha", "Hora": "Hora", "Transcripción": "Transcripción"})
    assert a.formatear_hora(lg.loc[0, "hora"]) == "08:30"
    assert a.formatear_hora(lg.loc[1, "hora"]) == ""
    assert a.formatear_hora(510.0) == "08:30" and a.formatear_hora(float("nan")) == ""


def test_hora_de_columnas_suma_al_puntaje():
    df = pd.DataFrame({"Descripción": ["Falla señal 14"], "Fecha": ["04/03/2025"], "N° Novedad": [1],
                       "Hora": ["08:20"]})
    pcz = a.preparar_pcz(df, {"Descripción": "Descripción", "Fecha": "Fecha", "N° Novedad": "N° Novedad",
                              "Hora": "Hora"})
    assert list(pcz["hora"]) == [8 * 60 + 20]
    lg = pd.DataFrame({"fecha": [F(2025, 3, 4)] * 2, "hora": [8 * 60 + 15, 15 * 60],
                       "avisos": [["Falla señal 14"], ["Falla señal 14"]]})
    c = a.candidatos(pcz.iloc[0], a.tabla_avisos(lg))
    assert list(c["clave"]) == ["0:0", "1:0"]  # el de hora cercana primero
    assert c.iloc[0]["puntaje"] == c.iloc[1]["puntaje"] + 3 and "hora a 5 min" in c.iloc[0]["detalle"]


