"""Genera un Excel de prueba INVENTADO para el módulo Informes PowerPoint (etapa 2).

Uso:  python tests/generar_ejemplo_excel.py
Crea tests/ejemplos/Indicadores_ejemplo.xlsx, parecido a la planilla de indicadores: carátula con
el mes, una hoja con una tabla y un gráfico de columnas, otra con una torta y un treemap (gráfico
moderno, que openpyxl no sabe hacer: se agrega a mano al zip, como lo guarda Excel, envuelto en
mc:AlternateContent y dentro de un grupo) y una hoja oculta.
"""
import datetime as dt
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference

ARCHIVO = Path(__file__).parent / "ejemplos" / "Indicadores_ejemplo.xlsx"

TREEMAP_ANCLA = """<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<mc:Choice xmlns:cx1="http://schemas.microsoft.com/office/drawing/2015/9/8/chartex" Requires="cx1">
<xdr:twoCellAnchor><xdr:from><xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>20</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>13</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>38</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="20" name="Grupo 1"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr><xdr:grpSpPr/>
<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="21" name="Gráfico 7"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>
<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2014/chartex">
<cx:chart xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex" r:id="rIdTreemap"/></a:graphicData></a:graphic>
</xdr:graphicFrame></xdr:grpSp><xdr:clientData/></xdr:twoCellAnchor></mc:Choice>
<mc:Fallback><xdr:twoCellAnchor><xdr:from><xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>20</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>13</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>38</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="0" name=""/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/>
<xdr:txBody><a:bodyPr/><a:p><a:r><a:t>Este gráfico no está disponible en su versión de Excel.</a:t></a:r></a:p></xdr:txBody></xdr:sp>
<xdr:clientData/></xdr:twoCellAnchor></mc:Fallback></mc:AlternateContent>"""

TREEMAP = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cx:chartSpace xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex">
<cx:chartData><cx:data id="0"><cx:strDim type="cat"><cx:f>'Disponib.'!$A$2:$A$6</cx:f></cx:strDim>
<cx:numDim type="size"><cx:f>'Disponib.'!$B$2:$B$6</cx:f></cx:numDim></cx:data></cx:chartData>
<cx:chart><cx:title pos="t" align="ctr" overlay="0"><cx:tx><cx:txData><cx:v>Fallas por sistema</cx:v></cx:txData></cx:tx>
<cx:txPr><a:bodyPr/><a:p><a:r><a:t>Fallas por sistema</a:t></a:r></a:p></cx:txPr></cx:title>
<cx:plotArea><cx:plotAreaRegion><cx:series layoutId="treemap" uniqueId="{00000000-0000-0000-0000-000000000001}">
<cx:dataId val="0"/></cx:series></cx:plotAreaRegion></cx:plotArea></cx:chart></cx:chartSpace>"""


def _agregar_treemap(contenido: bytes) -> bytes:
    """Agrega el treemap al dibujo de la hoja «Disponib.» (drawing2.xml)."""
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in entrada.infolist():
            datos = entrada.read(item.filename)
            if item.filename == "xl/drawings/drawing2.xml":
                datos = datos.replace(b"</wsDr>", TREEMAP_ANCLA.encode() + b"</wsDr>")
            elif item.filename == "xl/drawings/_rels/drawing2.xml.rels":
                datos = datos.replace(b"</Relationships>", b'<Relationship Id="rIdTreemap" Type="http://schemas.'
                                      b'microsoft.com/office/2014/relationships/chartEx" Target="../charts/chartEx1.xml"/>'
                                      b"</Relationships>")
            elif item.filename == "[Content_Types].xml":
                datos = datos.replace(b"</Types>", b'<Override PartName="/xl/charts/chartEx1.xml" ContentType='
                                      b'"application/vnd.ms-office.chartex+xml"/></Types>')
            z.writestr(item, datos)
        z.writestr("xl/charts/chartEx1.xml", TREEMAP)
    return salida.getvalue()


def generar(ruta: Path = ARCHIVO) -> Path:
    wb = Workbook()
    car = wb.active
    car.title = "Carátula"
    car["B2"] = "Informe mensual de indicadores RAMS"
    car["C32"], car["D32"] = "Período:", dt.datetime(2026, 9, 1)
    car["D32"].number_format = "mmmm yyyy"
    car["C33"], car["D33"] = "MKBF (km):", 12500.456
    car["D33"].number_format = "#,##0.0"
    car["C34"], car["D34"] = "Disponibilidad:", 0.875
    car["D34"].number_format = "0.0%"

    conf = wb.create_sheet("Confiabilidad")
    conf["A4"] = "Eventos por mes"
    conf.append([])
    for col, titulo in enumerate(["Mes", "Fallas", "Km", "MKBF"], 1):
        conf.cell(5, col, titulo)
    for i, mes in enumerate(["Jul", "Ago", "Sep"]):
        conf.append([mes, 31 - 4 * i + 15 * (i == 2), 390000 + 1000 * i, None])
    for f in range(42, 6, -1):
        conf.cell(f, 14, "")  # la tabla llega hasta N41
    conf.cell(41, 14, "fin")
    conf.column_dimensions["A"].width = 18
    barras = BarChart()
    barras.type = "col"
    barras.title = "Fallas por mes"
    barras.add_data(Reference(conf, min_col=2, min_row=5, max_row=8), titles_from_data=True)
    barras.set_categories(Reference(conf, min_col=1, min_row=6, max_row=8))
    conf.add_chart(barras, "P5")

    disp = wb.create_sheet("Disponib.")
    disp.append(["Sistema", "Fallas", "Disp. %"])
    for fila in [["Señales", 14, 0.91], ["Barreras", 9, 0.88], ["Pasos a nivel", 7, 0.93],
                 ["Comunicaciones", 5, 0.97], ["Energía", 7, 0.95]]:
        disp.append(fila)
    torta = PieChart()
    torta.title = "Participación por sistema"
    torta.add_data(Reference(disp, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    torta.set_categories(Reference(disp, min_col=1, min_row=2, max_row=6))
    disp.add_chart(torta, "F2")

    oculta = wb.create_sheet("Datos")
    oculta.sheet_state = "hidden"
    oculta["A1"] = "datos auxiliares"

    salida = io.BytesIO()
    wb.save(salida)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_bytes(_agregar_treemap(salida.getvalue()))
    return ruta


if __name__ == "__main__":
    print(f"Creado: {generar()}")
