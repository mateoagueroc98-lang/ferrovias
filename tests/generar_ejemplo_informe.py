"""Genera un PowerPoint de prueba INVENTADO para el módulo Informes PowerPoint.

Uso:  python tests/generar_ejemplo_informe.py
Crea tests/ejemplos/Informe_ejemplo.pptx, parecido a un informe mensual: título con el mes partido
en varios runs (como lo guarda PowerPoint), capturas pegadas como imagen, un gráfico, un grupo
(cuadro de texto + imagen), una tabla, un marcador de imagen vacío y una forma decorativa sin texto.
"""
import io
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm, Pt

ARCHIVO = Path(__file__).parent / "ejemplos" / "Informe_ejemplo.pptx"
ROJO = RGBColor(0xD4, 0x19, 0x20)


def captura(ancho: int, alto: int, filas: int, color: str) -> io.BytesIO:
    """Imagen que parece una captura de un rango de Excel (grilla con encabezado de color)."""
    img = Image.new("RGB", (ancho, alto), "white")
    d = ImageDraw.Draw(img)
    alto_fila = alto // filas
    d.rectangle([0, 0, ancho, alto_fila], fill=color)
    for f in range(filas + 1):
        d.line([0, f * alto_fila, ancho, f * alto_fila], fill="#9A9A9A")
    for c in range(5):
        d.line([c * ancho // 4, 0, c * ancho // 4, alto], fill="#9A9A9A")
    salida = io.BytesIO()
    img.save(salida, "PNG")
    salida.seek(0)
    return salida


def generar(ruta: Path = ARCHIVO) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)  # 16:9

    # 1. Portada: el mes queda partido en dos runs ("Septi" + "embre 2026"), con el mismo formato
    s = prs.slides.add_slide(prs.slide_layouts[0])
    p = s.shapes.title.text_frame.paragraphs[0]
    for texto, negrita in [("Informe mensual de mantenimiento – ", False), ("Septi", True), ("embre 2026", True)]:
        r = p.add_run()
        r.text = texto
        r.font.bold = negrita
        if negrita:
            r.font.color.rgb = ROJO
    s.placeholders[1].text = "Ingeniería de Mantenimiento · RAMS"

    # 2. Confiabilidad: captura pegada como imagen, gráfico y un texto con el mes en el medio
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Confiabilidad"
    s.shapes.add_picture(captura(800, 400, 8, "#C4161D"), Cm(1.5), Cm(4.5), Cm(15), Cm(7.5))
    datos = CategoryChartData()
    datos.categories = ["Jul", "Ago", "Sep"]
    datos.add_series("Fallas", (31, 27, 42))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(17.5), Cm(4.5), Cm(15), Cm(7.5), datos)
    caja = s.shapes.add_textbox(Cm(1.5), Cm(13), Cm(31), Cm(1.5))
    p = caja.text_frame.paragraphs[0]
    for texto in ["Eventos registrados en ", "Septiembre", " 2026: 42 (MKBF 12.500 km)"]:
        r = p.add_run()
        r.text = texto
        r.font.size = Pt(16)
    # Forma decorativa sin texto: no se tiene que listar
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(18.5), Cm(33.867), Cm(0.55))

    # 3. Disponibilidad: grupo (cuadro de texto + captura) y una tabla
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Disponibilidad de flota"
    grupo = s.shapes.add_group_shape()
    grupo.shapes.add_textbox(Cm(1.5), Cm(4.5), Cm(14), Cm(1.2)).text_frame.text = "Locomotoras GT22"
    grupo.shapes.add_picture(captura(700, 350, 6, "#2F7D4A"), Cm(1.5), Cm(6), Cm(14), Cm(7))
    tabla = s.shapes.add_table(3, 3, Cm(17.5), Cm(4.5), Cm(15), Cm(4.5)).table
    for f, fila in enumerate([["Flota", "Disp. %", "Meta %"], ["GT22", "87,5", "90"], ["Coches", "", "92"]]):
        for c, valor in enumerate(fila):
            tabla.cell(f, c).text = valor

    # 4. Marcador de imagen vacío (diseño "Imagen con descripción") y un pie de fuente
    s = prs.slides.add_slide(prs.slide_layouts[8])
    s.shapes.title.text = "Eventos destacados"
    s.placeholders[2].text = "Detalle de los eventos de Septiembre 2026."
    s.shapes.add_textbox(Cm(1.5), Cm(17.5), Cm(20), Cm(1)).text_frame.text = "Fuente: planilla Indicadores RAMS"

    ruta.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ruta))
    return ruta


if __name__ == "__main__":
    print(f"Creado: {generar()}")
