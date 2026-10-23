"""Tests del módulo Concatenación (correr con:  python -m pytest)."""
import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from src import concatenacion as ln
from tests.generar_ejemplo_concatenacion import ARCHIVO, generar

T = dt.time
CONECTORES = ln.CONECTORES_DEFECTO
PREVIA = "Textbox5,Textbox12\r\n*Se muestran las novedades pendientes\r\n\r\n"
ENCABEZADO = "numero,Descripción_Estado,fecha,hora,descrip\r\n"


@pytest.fixture(scope="module")
def ejemplo() -> bytes:
    generar()
    return ARCHIVO.read_bytes()


def csv(*renglones: str, encoding: str = "cp1252") -> bytes:
    return (PREVIA + ENCABEZADO + "".join(r + "\r\n" for r in renglones)).encode(encoding)


def por_numero(resultado: ln.Resultado) -> dict:
    return {f["Nro novedad"]: f for f in resultado.tabla.to_dict("records")}


# ---------------------------------------------------------- normalización
def test_normalizar():
    assert ln.normalizar("DERIVACIÓN") == "derivacion"
    assert ln.normalizar("SEÑAL") == "senal"
    assert ln.normalizar("NORMAL.-") == "normal"
    assert ln.normalizar("MOVIL 37  DE") == "movil 37 de"
    assert ln.normalizar('"B" C/A.-') == "b c a"


@pytest.mark.parametrize("texto, clave, coincide", [
    ("MEC. SILVA DA NORMAL.-", "NORMAL", True),
    ("FUNCIONAMIENTO ANORMAL", "NORMAL", False),     # palabra completa
    ("SE ESPERA NORMALIZADO", "NORMAL", False),
    ("NORMAL EL SECTOR.-", "normal el sector", True),
    ("NORMAL EL SECTOR.-", "normal sector", True),   # conectores fuera en los dos lados
    ("INSP. GALVAN RECIBE VIA ASC.", "recibe via", True),
    ("VIA LIBRE, RECIBE CUADRO", "recibe via", False),  # en otro orden
    ("RECIBE CUADRO", "recibe via", False),
    ("SEÑAL 23 NO ACLARA", "23 no aclara", True),  # "no" nunca es conector
    ("SEÑAL 23 ACLARA", "senal 23 no aclara", False),
    ("DERIVACIÓN NORMALIZADA", "derivacion normalizada", True),
])
def test_coincidencia(texto, clave, coincide):
    conectores, _ = ln.preparar_conectores(CONECTORES)
    claves, _ = ln.preparar_claves([clave], conectores)
    assert ln.contiene(ln.tokens(texto, conectores), claves[0]) is coincide


def test_no_y_sin_nunca_son_conectores():
    conectores, rechazados = ln.preparar_conectores(CONECTORES + ["no", "SIN", "tras"])
    assert rechazados == ["no", "SIN"] and "tras" in conectores


def test_clave_vacia_se_ignora():
    conectores, _ = ln.preparar_conectores(CONECTORES)
    claves, ignoradas = ln.preparar_claves(["de la", "NORMAL", " .- ", "normal"], conectores)
    assert claves == [["normal"]] and ignoradas == ["de la", " .- "]


# --------------------------------------------- ejemplos del documento del módulo
ESPERADO = {
    799: (T(3, 16), None), 800: (T(0, 43), None), 805: (T(2, 18), T(2, 18)),
    807: (T(4, 48), T(7, 45)), 808: (T(4, 49), T(5, 50)), 813: (T(6, 41), T(7, 52)),
    817: (T(9, 6), None), 820: (T(10, 14), T(10, 58)), 821: (T(10, 31), T(10, 41)),
    824: (T(10, 45), T(11, 51)),
}


def test_tabla_de_ejemplos(ejemplo):
    filas = por_numero(ln.procesar(ejemplo, ["NORMAL"], CONECTORES))
    for nro, (inicio, fin) in ESPERADO.items():
        assert filas[nro]["Fecha"] == dt.date(2026, 10, 21)
        assert (filas[nro]["Hora inicio"], filas[nro]["Hora fin"]) == (inicio, fin), nro


def test_clave_recibe_via(ejemplo):
    filas = por_numero(ln.procesar(ejemplo, ["NORMAL", "recibe via"], CONECTORES))
    assert filas[799]["Hora fin"] == T(3, 16)   # coincide la primera fila
    assert filas[802]["Hora fin"] == T(1, 55)   # varias: la última
    assert filas[804]["Hora fin"] is None       # "RECIBE CUADRO" no es "recibe via"


def test_una_fila_por_novedad_en_orden(ejemplo):
    r = ln.procesar(ejemplo, ["NORMAL"], CONECTORES)
    assert list(r.tabla["Nro novedad"]) == [790, 799, 800, 802, 803, 804, 805, 807, 808,
                                             813, 817, 820, 821, 824]
    assert r.renglones == 30 and r.encoding == "Windows-1252"


def test_detalle(ejemplo):
    filas = por_numero(ln.procesar(ejemplo, ["NORMAL"], CONECTORES))
    assert filas[808]["Detalle"] == (
        "AUX. BOGADO EST. G.BOURG INF. QUE SEÑAL 23 NO ACLARA PARA T.3012 LOC.\n"
        "E703, COND. CAPPELLETTI INF. QUE ESTA APAGADA.-\n"
        "MEC. ENRIQUEZ C/A.-\n"
        "MEC. SILVA DA NORMAL.-")
    # comas internas entre comillas y comillas internas, completas
    assert filas[803]["Detalle"].startswith("MITAD DE PLATAFORMA DESC., EL FUE ADVERTIDO")
    assert filas[817]["Detalle"] == 'HABILITAN ESTACIONES "B" CON CUSTODIA POR FALTA DE PERSONAL.-'
    # tal cual: sin unir palabras cortadas con guion
    assert filas[800]["Detalle"].startswith("T.3001 DETE-\nNIDO EN SEÑAL 41")


def test_pendiente_anterior_y_numero_separado(ejemplo):
    r = ln.procesar(ejemplo, ["NORMAL"], CONECTORES)
    fila = por_numero(r)[790]
    assert fila["Fecha"] == dt.date(2026, 10, 20)
    assert fila["Detalle"].endswith("\nSIGUE SIN COMPROBACION.-")
    assert any("separados" in a and "790" in a for a in r.avisos)


def test_estado_no_cuenta_como_palabra_clave():
    r = ln.procesar(csv("1,NORMAL,21/10/2026,08:00,SIN NOVEDAD.-"), ["NORMAL"], CONECTORES)
    assert por_numero(r)[1]["Hora fin"] is None


# ------------------------------------------------------------- lectura
def test_encoding_automatico_utf8_y_cp1252():
    for enc in ("cp1252", "utf-8", "utf-8-sig"):
        r = ln.procesar(csv("5,NORMAL,21/10/2026,08:00,IBAÑEZ DERIVACIÓN", encoding=enc),
                        ["NORMAL"], CONECTORES)
        assert por_numero(r)[5]["Detalle"] == "IBAÑEZ DERIVACIÓN"


def test_encoding_utf8_forzado_con_archivo_windows():
    with pytest.raises(ln.ErrorArchivo, match="UTF-8"):
        ln.procesar(csv("5,NORMAL,21/10/2026,08:00,IBAÑEZ"), [], CONECTORES, "utf-8-sig")


def test_encabezado_distinto_se_rechaza():
    malo = (PREVIA + "nro,estado,fecha,hora,texto\r\n1,NORMAL,21/10/2026,08:00,X\r\n").encode()
    with pytest.raises(ln.ErrorArchivo, match="encabezado"):
        ln.procesar(malo, ["NORMAL"], CONECTORES)


def test_encabezado_sin_filas_previas():
    r = ln.procesar((ENCABEZADO + "1,NORMAL,21/10/2026,08:00,X\r\n").encode(), [], CONECTORES)
    assert len(r.tabla) == 1


@pytest.mark.parametrize("contenido", [b"", b"\r\n", csv()])
def test_archivo_vacio(contenido):
    r = ln.procesar(contenido, ["NORMAL"], CONECTORES)
    assert r.tabla.empty and any("no tiene novedades" in a for a in r.avisos)
    hoja = load_workbook(io.BytesIO(ln.exportar_excel(r.tabla))).active
    assert [c.value for c in hoja[1]] == ln.COLUMNAS and hoja.max_row == 1


def test_fecha_y_hora_invalidas():
    r = ln.procesar(csv("1,NORMAL,31/02/2026,25:10,X", "1,NORMAL,31/02/2026,8.30,NORMAL"),
                    ["NORMAL"], CONECTORES)
    fila = por_numero(r)[1]
    assert fila["Fecha"] is None and fila["Hora inicio"] is None and fila["Hora fin"] is None
    assert any("Fecha" in a and "fila 5" in a for a in r.avisos)
    assert any("Hora" in a and "fila 5" in a and "fila 6" in a for a in r.avisos)


def test_numero_invalido_se_saltea():
    r = ln.procesar(csv("abc,NORMAL,21/10/2026,08:00,X", "2,NORMAL,21/10/2026,08:00,Y"),
                    [], CONECTORES)
    assert list(r.tabla["Nro novedad"]) == [2] and any("filas 5." in a for a in r.avisos)


def test_coma_sin_comillas_se_junta():
    r = ln.procesar(csv("1,NORMAL,21/10/2026,08:00,UNO, DOS"), [], CONECTORES)
    assert por_numero(r)[1]["Detalle"] == "UNO, DOS"


def test_cruce_de_medianoche_no_se_corrige():
    r = ln.procesar(csv("1,NORMAL,21/10/2026,23:50,X", "1,NORMAL,21/10/2026,00:20,NORMAL.-"),
                    ["NORMAL"], CONECTORES)
    assert por_numero(r)[1]["Hora fin"] == T(0, 20)


def test_detalle_muy_largo_se_trunca():
    r = ln.procesar(csv(*["1,NORMAL,21/10/2026,08:00," + "X" * 1000] * 40), [], CONECTORES)
    assert len(por_numero(r)[1]["Detalle"]) == ln.MAX_CELDA
    assert any("recortó" in a for a in r.avisos)


# ------------------------------------------------------------ exportación
def test_excel(ejemplo):
    r = ln.procesar(ejemplo, ["NORMAL"], CONECTORES)
    hoja = load_workbook(io.BytesIO(ln.exportar_excel(r.tabla))).active
    assert [c.value for c in hoja[1]] == ln.COLUMNAS and all(c.font.bold for c in hoja[1])
    assert hoja.auto_filter.ref == f"A1:E{len(r.tabla) + 1}"
    fila = [c for c in hoja[10]]  # 808
    assert fila[0].value == 808 and isinstance(fila[0].value, int)
    assert fila[1].value == dt.datetime(2026, 10, 21) and fila[1].number_format == "dd/mm/yyyy"
    assert fila[2].value == T(4, 49) and fila[2].number_format == "hh:mm"
    assert fila[3].value == T(5, 50)
    assert fila[4].alignment.wrap_text and fila[4].value.count("\n") == 3
    assert hoja.cell(3, 4).value is None  # 799 sin hora fin: celda vacía


def test_excel_formato_tabla(ejemplo):
    r = ln.procesar(ejemplo, ["NORMAL"], CONECTORES)
    hoja = load_workbook(io.BytesIO(ln.exportar_excel(r.tabla, ln.FORMATO_TABLA, "Tabla"))).active
    assert hoja.title == "Tabla"
    assert [c.value for c in hoja[1]] == [
        "Fecha", "Mes", "N° Novedad", "Sistema", "Subsistema", "Parte", "Elemento",
        "Nombre correcto", "Descripción", "Zona", "VER", "Falla", "Comienzo de falla", "Habilitación"]
    assert hoja.auto_filter.ref == f"A1:N{len(r.tabla) + 1}"
    fila = hoja[10]  # 808
    assert fila[0].value == dt.datetime(2026, 10, 21) and fila[0].number_format == "dd/mm/yyyy"
    assert fila[2].value == 808
    assert fila[8].value.startswith("AUX. BOGADO") and fila[8].alignment.wrap_text
    assert fila[12].value == T(4, 49) and fila[12].number_format == "hh:mm"
    assert fila[13].value == T(5, 50)
    # columnas vacías: sin ningún valor, para que «Saltar blancos» no pise las fórmulas
    for i in (1, 3, 4, 5, 6, 7, 9, 10, 11):
        assert fila[i].value is None
    assert hoja.cell(3, 14).value is None  # 799 sin hora fin
