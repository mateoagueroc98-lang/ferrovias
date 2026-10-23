"""Tests del módulo Informes PowerPoint, etapa 1: plantillas (correr con:  python -m pytest)."""
import io
import json

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

from src.informes_pptx import esquema, office, reemplazo
from src.informes_pptx import plantillas as pl
from tests.generar_ejemplo_informe import ARCHIVO, generar


@pytest.fixture(scope="module")
def ejemplo() -> bytes:
    generar()
    return ARCHIVO.read_bytes()


@pytest.fixture
def plantilla(ejemplo, tmp_path) -> dict:
    return pl.crear("Informe mensual RAMS", "Informe_ejemplo.pptx", ejemplo, tmp_path)


def cuadro(*runs) -> object:
    """Un cuadro de texto con un párrafo formado por estos runs: (texto, negrita, tamaño)."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(2)).text_frame
    p = tf.paragraphs[0]
    for texto, negrita, tamano in runs:
        r = p.add_run()
        r.text = texto
        r.font.bold = negrita
        r.font.size = Pt(tamano)
    return tf


def runs(tf) -> list[tuple]:
    return [(r.text, r.font.bold, r.font.size.pt if r.font.size else None)
            for p in tf.paragraphs for r in p.runs]


def textos_de(diap: dict) -> list[str]:
    return [t["contenido"] for t in diap["textos"]]


# ------------------------------------------------------------------ reemplazo
def test_fragmento_partido_en_varios_runs_conserva_el_formato():
    tf = cuadro(("Informe – ", False, 20), ("Septi", True, 28), ("embre 2026", True, 28))
    assert reemplazo.reemplazar_fragmento(tf, "Septiembre 2026", "Octubre 2026") == 1
    # El texto nuevo queda en el run donde empezaba ("Septi"), con su negrita y tamaño
    assert runs(tf) == [("Informe – ", False, 20), ("Octubre 2026", True, 28)]


def test_fragmento_dentro_de_un_run_y_entre_dos_formatos():
    tf = cuadro(("Mes: Septiembre", False, 12), (" 2026 final", True, 14))
    reemplazo.reemplazar_fragmento(tf, "Septiembre 2026", "Octubre 2026")
    assert runs(tf) == [("Mes: Octubre 2026", False, 12), (" final", True, 14)]


def test_fragmento_varias_veces_y_ninguna():
    tf = cuadro(("Sep a Sep", False, 12))
    assert reemplazo.reemplazar_fragmento(tf, "Sep", "Oct") == 2
    assert reemplazo.texto_cuadro(tf) == "Oct a Oct"
    assert reemplazo.reemplazar_fragmento(tf, "Nov", "Dic") == 0
    assert reemplazo.texto_cuadro(tf) == "Oct a Oct"


def test_fragmento_con_texto_nuevo_vacio_borra_runs():
    tf = cuadro(("A", False, 12), ("B", True, 12), ("C", False, 12))
    reemplazo.reemplazar_fragmento(tf, "B", "")
    assert runs(tf) == [("A", False, 12), ("C", False, 12)]


def test_reemplazar_todo_usa_el_formato_del_primer_run_y_separa_renglones():
    tf = cuadro(("Título ", True, 30), ("viejo", False, 10))
    tf.add_paragraph().text = "segundo párrafo"
    reemplazo.reemplazar_todo(tf, "Nuevo\nOtro renglón")
    assert runs(tf) == [("Nuevo", True, 30), ("Otro renglón", True, 30)]
    assert len(tf.paragraphs) == 2


def test_reemplazar_todo_en_cuadro_vacio_toma_el_formato_de_fin_de_parrafo():
    prs = Presentation()
    tf = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Cm(5), Cm(2)).text_frame
    fin = tf.paragraphs[0]._p.get_or_add_endParaRPr()
    fin.set("sz", "2400")
    fin.set("b", "1")
    reemplazo.reemplazar_todo(tf, "Hola")
    assert runs(tf) == [("Hola", True, 24)]


def test_salto_de_linea_cuenta_como_renglon():
    tf = cuadro(("uno", False, 12))
    p = tf.paragraphs[0]
    p.add_line_break()
    r = p.add_run()
    r.text = "dos"
    assert reemplazo.texto_cuadro(tf) == "uno\ndos"
    assert reemplazo.reemplazar_fragmento(tf, "dos", "tres") == 1
    assert reemplazo.texto_cuadro(tf) == "uno\ntres"


# -------------------------------------------------------------------- detección
def test_detecta_textos_e_imagenes_de_cada_diapositiva(plantilla):
    d1, d2, d3, d4 = plantilla["diapositivas"]
    assert textos_de(d1) == ["Informe mensual de mantenimiento – Septiembre 2026", "Ingeniería de Mantenimiento · RAMS"]
    assert d1["imagenes"] == []

    # Imagen pegada y gráfico; la forma decorativa sin texto no aparece
    assert textos_de(d2) == ["Confiabilidad", "Eventos registrados en Septiembre 2026: 42 (MKBF 12.500 km)"]
    assert [i["clase"] for i in d2["imagenes"]] == ["Imagen", "Gráfico"]

    # Texto e imagen dentro de un grupo, celdas de tabla (la vacía no)
    assert "Locomotoras GT22" in textos_de(d3)
    celdas = [t for t in d3["textos"] if t["celda"]]
    assert len(celdas) == 8
    assert {"Flota", "87,5", "92"} <= {t["contenido"] for t in celdas}
    assert next(t for t in d3["textos"] if t["contenido"] == "87,5")["celda"] == [2, 2]
    assert d3["imagenes"][0]["clase"] == "Imagen (en un grupo)"

    # Marcador de imagen vacío
    assert [i["clase"] for i in d4["imagenes"]] == ["Marcador de imagen vacío"]
    assert "Fuente: planilla Indicadores RAMS" in textos_de(d4)


def test_posiciones_y_tamano(plantilla):
    tam = plantilla["tamano_diapositiva"]
    assert (tam["ancho_cm"], tam["alto_cm"]) == (33.87, 19.05)
    img = plantilla["diapositivas"][1]["imagenes"][0]
    assert img["caja"] == {"x": Cm(1.5), "y": Cm(4.5), "ancho": Cm(15), "alto": Cm(7.5)}
    celda = next(t for t in plantilla["diapositivas"][2]["textos"] if t["contenido"] == "87,5")
    assert celda["caja"]["x"] == Cm(17.5) + Cm(5)  # segunda columna de la tabla


def test_posicion_dentro_de_un_grupo_escalado():
    """Si el grupo se agrandó, las coordenadas de adentro se pasan a las de la diapositiva."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    g = s.shapes.add_group_shape()
    g.shapes.add_textbox(Cm(2), Cm(2), Cm(4), Cm(2)).text_frame.text = "adentro"
    xfrm = g._element.grpSpPr.find(qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("x", str(Cm(10)))   # el grupo se movió a x = 10 cm
    xfrm.find(qn("a:ext")).set("cx", str(Cm(8)))   # y se estiró al doble de ancho
    caja = pl.analizar(prs)["diapositivas"][0]["textos"][0]["caja"]
    assert (caja["x"], caja["ancho"]) == (Cm(10), Cm(8))


def test_ids_repetidos_se_renumeran():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_textbox(0, 0, Cm(1), Cm(1))
    b = s.shapes.add_textbox(0, 0, Cm(1), Cm(1))
    b._element.nvSpPr.cNvPr.set("id", str(a.shape_id))
    assert pl.asegurar_ids_unicos(prs) == 1
    assert a.shape_id != b.shape_id


# --------------------------------------------------------------- en disco
def test_crear_guarda_base_y_json_y_nombra_los_lugares_de_imagen(plantilla, tmp_path):
    c = pl.carpeta(plantilla["id"], tmp_path)
    assert (c / "base.pptx").exists()
    datos = json.loads((c / "plantilla.json").read_text(encoding="utf-8"))
    assert datos["nombre"] == "Informe mensual RAMS" and datos["creada"]
    prs = Presentation(str(c / "base.pptx"))
    for d in plantilla["diapositivas"]:
        for img in d["imagenes"]:
            assert pl.lugar_de_imagen(prs, d["numero"], img).name == img["id"]


def test_lugar_de_imagen_se_encuentra_aunque_cambie_el_orden(plantilla, tmp_path):
    base = pl.carpeta(plantilla["id"], tmp_path) / "base.pptx"
    prs = Presentation(str(base))
    arbol = prs.slides[1].shapes._spTree
    grafico = pl.buscar_shape(prs.slides[1], nombre="img_02_02")._element
    arbol.remove(grafico)
    arbol.insert(2, grafico)  # el gráfico pasa a estar antes que la imagen
    grafico.xpath("./p:nvGraphicFramePr/p:cNvPr")[0].set("id", "99")  # y hasta cambia de número
    img = plantilla["diapositivas"][1]["imagenes"][1]
    assert pl.lugar_de_imagen(prs, 2, img).has_chart


def marcar(plantilla, diap, contenido, **cambios) -> dict:
    t = next(t for t in plantilla["diapositivas"][diap - 1]["textos"] if t["contenido"] == contenido)
    t.update(tipo=pl.VARIABLE, **cambios)
    return t


def test_guardar_valida(plantilla, tmp_path):
    t = marcar(plantilla, 1, "Informe mensual de mantenimiento – Septiembre 2026",
               reemplazo=pl.FRAGMENTO, fragmento="Octubre", variable_nombre="", variable_id="Mal id")
    with pytest.raises(pl.ErrorPlantilla) as e:
        pl.guardar(plantilla, tmp_path)
    assert "falta el nombre" in str(e.value) and "identificador" in str(e.value) and "no aparece" in str(e.value)
    t.update(fragmento="Septiembre 2026", variable_nombre="Mes del informe", variable_id="mes")
    guardada = pl.guardar(plantilla, tmp_path)
    assert pl.cargar(plantilla["id"], tmp_path)["diapositivas"][0]["textos"][0]["fragmento"] == "Septiembre 2026"
    assert guardada["modificada"] >= plantilla["creada"]


def test_mismo_identificador_con_nombres_distintos_no_se_guarda(plantilla, tmp_path):
    marcar(plantilla, 1, "Informe mensual de mantenimiento – Septiembre 2026", variable_nombre="Mes", variable_id="mes")
    marcar(plantilla, 4, "Fuente: planilla Indicadores RAMS", variable_nombre="Fuente", variable_id="mes")
    assert any("nombres distintos" in e for e in pl.validar(plantilla))


def test_variable_compartida_y_prueba_de_textos(plantilla, tmp_path):
    marcar(plantilla, 1, "Informe mensual de mantenimiento – Septiembre 2026", reemplazo=pl.FRAGMENTO,
           fragmento="Septiembre 2026", variable_nombre="Mes", variable_id="mes")
    marcar(plantilla, 2, "Eventos registrados en Septiembre 2026: 42 (MKBF 12.500 km)", reemplazo=pl.FRAGMENTO,
           fragmento="Septiembre 2026", variable_nombre="Mes", variable_id="mes")
    marcar(plantilla, 4, "Fuente: planilla Indicadores RAMS", variable_nombre="Fuente", variable_id="fuente")
    assert pl.validar(plantilla) == []
    assert pl.variables(plantilla)["mes"]["usos"] == [(1, plantilla["diapositivas"][0]["textos"][0]["id"]),
                                                       (2, plantilla["diapositivas"][1]["textos"][1]["id"])]
    prs = Presentation(io.BytesIO(pl.prueba_textos(plantilla, tmp_path)))
    titulo = prs.slides[0].shapes.title.text_frame
    assert reemplazo.texto_cuadro(titulo) == "Informe mensual de mantenimiento – «Mes»"
    r = titulo.paragraphs[0].runs[-1]
    assert r.font.bold and r.font.color.rgb == RGBColor(0xD4, 0x19, 0x20)  # conserva negrita y color
    assert reemplazo.texto_cuadro(prs.slides[1].shapes[3].text_frame).startswith("Eventos registrados en «Mes»:")
    assert reemplazo.texto_cuadro(prs.slides[3].shapes[-1].text_frame) == "«Fuente»"


def test_listar_duplicar_eliminar(plantilla, tmp_path):
    copia = pl.duplicar(plantilla["id"], tmp_path)
    assert copia["nombre"] == "Informe mensual RAMS (copia)" and copia["id"] != plantilla["id"]
    assert (pl.carpeta(copia["id"], tmp_path) / "base.pptx").exists()
    (tmp_path / "plantillas" / "rota").mkdir()
    lista, rotas = pl.listar(tmp_path)
    assert {p["id"] for p in lista} == {plantilla["id"], copia["id"]} and rotas == ["rota"]
    pl.eliminar(plantilla["id"], tmp_path)
    assert [p["id"] for p in pl.listar(tmp_path)[0]] == [copia["id"]]


def test_archivo_que_no_es_powerpoint(tmp_path):
    with pytest.raises(pl.ErrorPlantilla):
        pl.crear("x", "x.pptx", b"no es un pptx", tmp_path)


def test_slug():
    assert pl.slug("Mes del título") == "mes_del_titulo"
    assert pl.slug("  Disponibilidad – GT22 %") == "disponibilidad_gt22"


# -------------------------------------------------------- esquema y office
def test_esquema_dibuja_un_recuadro_por_texto_e_imagen(plantilla):
    d = plantilla["diapositivas"][1]
    svg = esquema.svg(d, plantilla["tamano_diapositiva"])
    assert svg.startswith("<svg") and ">T1<" in svg and ">T2<" in svg and ">I2<" in svg


def test_office_simulado_no_frena(monkeypatch, tmp_path):
    monkeypatch.setenv("INFORMES_OFFICE", "simulado")
    of = office.obtener()
    assert not of.hace_miniaturas and of.miniaturas(tmp_path / "x.pptx", tmp_path) == []
