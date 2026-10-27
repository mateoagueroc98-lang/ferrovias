"""Tests del módulo Transcripciones del LG (correr con:  python -m pytest)."""
import datetime as dt
import io
import json
import shutil
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image

from src import transcripcion_lg as tl
from tests.generar_ejemplo_transcripcion import ARCHIVO, PAGINAS, generar


def respuesta(filas, continuacion="", cortado=False) -> str:
    return json.dumps({"continuacion_del_anterior": continuacion, "filas": filas,
                       "cortado_al_final": cortado, "dudosos": 0, "ilegibles": 0})


def fila(folio="folio 18", fecha="05/10/2026", hora="06:00", texto="AVISO"):
    return {"folio": folio, "fecha": fecha, "hora": hora, "transcripcion": texto}


def tanda_ok(paginas, texto_json) -> tl.Tanda:
    return tl.Tanda(paginas, estado="ok", resultado=tl.validar_respuesta(texto_json))


# ------------------------------------------------------------ validación JSON
def test_validar_respuesta_correcta():
    r = tl.validar_respuesta(respuesta([fila(texto="aviso 1\naviso 2")], cortado=True))
    assert r["cortado_al_final"] is True
    assert r["continuacion_del_anterior"] == ""
    assert r["filas"] == [{"folio": "folio 18", "fecha": "05/10/2026", "hora": "06:00",
                           "transcripcion": "aviso 1\naviso 2", "revisar": []}]


@pytest.mark.parametrize("envoltorio", [
    "```json\n{}\n```",
    "Acá va el JSON:\n{}\nListo.",
    "  {}  ",
])
def test_validar_respuesta_tolera_envoltorios(envoltorio):
    r = tl.validar_respuesta(envoltorio.replace("{}", respuesta([fila()])))
    assert len(r["filas"]) == 1


@pytest.mark.parametrize("cruda, mensaje", [
    ("no es json", "no es JSON"),
    ('{"filas": [ {"folio": "folio 1", ', "mal formado"),
    ("[1, 2]", "no es un objeto"),
    ('{"continuacion_del_anterior": ""}', "filas"),
    ('{"filas": "x"}', "filas"),
    ('{"filas": [{"folio": "folio 1", "fecha": "01/01/2026", "hora": ""}]}', "transcripcion"),
    ('{"filas": ["texto suelto"]}', "fila 1"),
    ('{"filas": [], "continuacion_del_anterior": ""}', "ninguna fila"),
    ('{"filas": [], "continuacion_del_anterior": 5}', "texto"),
    ('{"filas": [{"transcripcion": "x"}], "cortado_al_final": "si"}', "true o false"),
])
def test_validar_respuesta_mal_formada(cruda, mensaje):
    with pytest.raises(tl.ErrorRespuesta, match=mensaje):
        tl.validar_respuesta(cruda)


def test_validar_respuesta_normaliza_y_marca():
    r = tl.validar_respuesta(respuesta([
        fila(fecha="5/10/26", hora="6:05"),
        fila(fecha="", hora=""),
        fila(fecha="31/02/2026", hora="25:00"),
        {"transcripcion": "sin nada más", "hora": None},
    ]))
    f = r["filas"]
    assert (f[0]["fecha"], f[0]["hora"], f[0]["revisar"]) == ("05/10/2026", "06:05", [])
    assert (f[1]["fecha"], f[1]["hora"], f[1]["revisar"]) == ("", "", ["fecha vacía"])
    assert f[2]["revisar"] == ["fecha no reconocida", "hora no reconocida"]
    assert f[2]["fecha"] == "31/02/2026"  # se deja como vino para corregirlo a mano
    assert f[3]["folio"] == "folio [ilegible]"


def test_continuacion_ausente_es_vacia():
    r = tl.validar_respuesta('{"filas": [{"folio": "folio 2", "fecha": "01/01/2026", '
                             '"hora": "", "transcripcion": "x"}]}')
    assert r["continuacion_del_anterior"] == "" and r["cortado_al_final"] is False


@pytest.mark.parametrize("texto, esperado", [
    ("05/10/2026", "05/10/2026"), ("5/1/2026", "05/01/2026"), ("15.10.2026", "15/10/2026"),
    ("15-10-26", "15/10/2026"), ("2026-10-15", None), ("10/15/2026", None), ("", None),
])
def test_normalizar_fecha(texto, esperado):
    assert tl.normalizar_fecha(texto) == esperado


@pytest.mark.parametrize("texto, esperado", [
    ("06:00", "06:00"), ("6:5", None), ("6.30", "06:30"), ("0630", "06:30"), ("6 hs", "06:00"),
    ("23:59", "23:59"), ("24:00", None), ("de noche", None),
])
def test_normalizar_hora(texto, esperado):
    assert tl.normalizar_hora(texto) == esperado


# --------------------------------------------------------------------- folios
@pytest.mark.parametrize("folio, siguiente, esperado", [
    ("folio 18", "folio 19", "folio 18-19"),
    ("folio 17-18", "folio 19", "folio 17-19"),
    ("folio 18", "folio 18", "folio 18"),
    ("folio 18-19", "folio 19", "folio 18-19"),
    ("folio 99", "folio 01 (libro nuevo)", "folio 99-01 (libro nuevo)"),
    ("folio 01 (libro nuevo)", "folio 02", "folio 01-02 (libro nuevo)"),
    ("folio [ilegible]", "folio 19", "folio [ilegible]"),
])
def test_extender_folio(folio, siguiente, esperado):
    assert tl.extender_folio(folio, siguiente) == esperado


def test_folio_siguiente_respeta_ceros():
    assert tl.folio_siguiente("folio 08") == "folio 09"
    assert tl.folio_siguiente("folio 17-18") == "folio 19"


# ------------------------------------------------------------ unión de tandas
def test_union_con_continuacion():
    t1 = tanda_ok([0, 1, 2], respuesta([fila("folio 17", texto="A"),
                                         fila("folio 18", hora="14:00", texto="B empieza")], cortado=True))
    t2 = tanda_ok([3, 4, 5], respuesta([fila("folio 19", hora="22:00", texto="C")],
                                        continuacion="y B termina"))
    filas = tl.unir_tandas([t1, t2])
    assert [f.transcripcion for f in filas] == ["A", "B empieza y B termina", "C"]
    assert [f.folio for f in filas] == ["folio 17", "folio 18-19", "folio 19"]
    assert [f.tanda for f in filas] == [0, 0, 1]
    assert not any(f.revisar for f in filas)


def test_union_tanda_que_es_solo_continuacion():
    t1 = tanda_ok([0], respuesta([fila("folio 18", texto="B empieza")], cortado=True))
    t2 = tanda_ok([1], respuesta([], continuacion="sigue toda la página", cortado=True))
    t3 = tanda_ok([2], respuesta([fila("folio 21", texto="D")], continuacion="y termina"))
    filas = tl.unir_tandas([t1, t2, t3])
    assert [f.transcripcion for f in filas] == ["B empieza sigue toda la página y termina", "D"]
    assert filas[0].folio == "folio 18-21"


def test_union_cortado_sin_continuacion_se_marca():
    t1 = tanda_ok([0], respuesta([fila(texto="A cortado")], cortado=True))
    t2 = tanda_ok([1], respuesta([fila("folio 19", texto="B")]))
    filas = tl.unir_tandas([t1, t2])
    assert filas[0].revisar == [tl.REVISAR_CORTADO]
    assert filas[1].revisar == []


def test_union_cortado_al_final_del_pdf():
    filas = tl.unir_tandas([tanda_ok([0], respuesta([fila()], cortado=True))])
    assert filas[0].revisar == ["quedó cortado al final del PDF"]


def test_union_tanda_con_error_en_el_medio():
    t1 = tanda_ok([0], respuesta([fila("folio 18", texto="A")], cortado=True))
    t2 = tl.Tanda([1], estado="error", error="JSON mal formado")
    t3 = tanda_ok([2], respuesta([fila("folio 20", fecha="06/10/2026", texto="C")], continuacion="resto"))
    filas = tl.unir_tandas([t1, t2, t3])
    # A quedó cortado (la tanda siguiente falló) y la continuación de la 3 no se pega a A.
    assert filas[0].revisar == [tl.REVISAR_CORTADO]
    assert filas[1].transcripcion == "resto" and filas[1].revisar == [tl.REVISAR_SIN_ANTERIOR]
    assert filas[1].fecha == "05/10/2026" and filas[1].folio == "folio 20"
    assert filas[2].transcripcion == "C"


def test_contexto_para():
    t1 = tanda_ok([0, 1, 2], respuesta([fila(fecha="05/10/2026", texto="A"),
                                         fila(fecha="", texto="B")], cortado=True))
    t2 = tl.Tanda([3, 4, 5])
    assert tl.contexto_para([t1, t2], 0, "04/10/2026") == {
        "ultima_fecha": "04/10/2026", "ultimo_parrafo": "", "anterior_cortado": False}
    assert tl.contexto_para([t1, t2], 1, "04/10/2026") == {
        "ultima_fecha": "05/10/2026", "ultimo_parrafo": "B", "anterior_cortado": True}
    t1.estado = "error"
    assert tl.contexto_para([t1, t2], 1, "04/10/2026")["ultimo_parrafo"] == ""


def test_mensaje_usuario():
    contexto = {"ultima_fecha": "05/10/2026", "ultimo_parrafo": "B empieza", "anterior_cortado": True}
    contenido = tl.mensaje_usuario([b"img1", b"img2"], [3, 4], contexto)
    tipos = [c["type"] for c in contenido]
    assert tipos == ["text", "image", "text", "image", "text"]
    assert contenido[0]["text"] == "Página 4 del PDF:"
    assert contenido[1]["source"]["media_type"] == "image/jpeg"
    final = contenido[-1]["text"]
    assert "05/10/2026" in final and "B empieza" in final and "cortado" in final


# -------------------------------------------------------- llamada y reintento
def test_procesar_tanda_reintenta_una_vez():
    respuestas = iter(["{ roto", respuesta([fila()])])
    uso = [{"modelo": "claude-opus-5-5", "entrada": 10, "salida": 5, "cache_escritura": 0, "cache_lectura": 0}]
    t = tl.procesar_tanda(tl.Tanda([0]), lambda s, c: (next(respuestas), uso), [], [], {"x": 1})
    assert t.estado == "ok" and len(t.intentos) == 2
    assert t.intentos[0]["error"] and t.intentos[0]["cruda"] == "{ roto"
    assert t.intentos[1]["error"] == ""


def test_procesar_tanda_falla_dos_veces():
    t = tl.procesar_tanda(tl.Tanda([0]), lambda s, c: ("nada", []), [], [], {})
    assert t.estado == "error" and len(t.intentos) == 2 and "JSON" in t.error


def test_procesar_tanda_negativa_del_modelo():
    def llamar(s, c):
        raise tl.ErrorRespuesta("El modelo se negó a responder esta tanda.", "texto", [])
    t = tl.procesar_tanda(tl.Tanda([0]), llamar, [], [], {})
    assert t.estado == "error" and t.intentos[0]["cruda"] == "texto"


class ClienteFalso:
    """Imita client.beta.messages.stream(...) del SDK y guarda los parámetros recibidos."""

    def __init__(self, texto, stop_reason="end_turn"):
        self.texto, self.stop_reason, self.parametros = texto, stop_reason, None
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self.stream))

    def stream(self, **parametros):
        self.parametros = parametros
        cliente = self

        class Stream:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get_final_message(self):
                parte = SimpleNamespace(model="claude-opus-5-5", input_tokens=1000, output_tokens=500,
                                        cache_creation_input_tokens=200, cache_read_input_tokens=0)
                return SimpleNamespace(
                    model="claude-opus-5-5", stop_reason=cliente.stop_reason,
                    content=[SimpleNamespace(type="thinking"), SimpleNamespace(type="text", text=cliente.texto)],
                    usage=SimpleNamespace(iterations=[parte]))
        return Stream()


def test_llamar_api():
    cliente = ClienteFalso('{"filas": []}')
    system = tl.system_prompt("INSTR", "REF")
    texto, uso = tl.llamar_api(cliente, "claude-opus-5-5", system, [{"type": "text", "text": "x"}])
    assert texto == '{"filas": []}'
    assert uso == [{"modelo": "claude-opus-5-5", "entrada": 1000, "salida": 500,
                    "cache_escritura": 200, "cache_lectura": 0}]
    p = cliente.parametros
    assert p["model"] == "claude-opus-5-5" and p["system"][0]["text"] == "INSTR"
    assert "REF" in p["system"][1]["text"] and p["system"][1]["cache_control"]["type"] == "ephemeral"
    assert p["messages"][0]["role"] == "user"


def test_llamar_api_respuesta_cortada():
    with pytest.raises(tl.ErrorRespuesta, match="cortada") as e:
        tl.llamar_api(ClienteFalso("{ a medias", "max_tokens"), "claude-opus-5-5", [], [])
    assert e.value.cruda == "{ a medias" and e.value.uso


def test_costo():
    uso = [{"modelo": "claude-opus-5-5", "entrada": 1_000_000, "salida": 100_000,
            "cache_escritura": 0, "cache_lectura": 1_000_000}]
    assert tl.costo(uso) == pytest.approx(4 + 2 + 0.2)
    assert tl.estimar_costo(6, 3000, "claude-sonnet-5-5") < tl.estimar_costo(6, 3000, "claude-opus-5-5")


# -------------------------------------------------------------- PDF y tandas
def test_armar_tandas_siempre_de_a_3():
    assert tl.armar_tandas(7) == [[0, 1, 2], [3, 4, 5], [6]]
    assert tl.armar_tandas(3) == [[0, 1, 2]]
    assert tl.armar_tandas(0) == []


def test_paginas_a_imagenes_con_pdf_de_ejemplo():
    if not ARCHIVO.exists():
        generar()
    imagenes = tl.paginas_a_imagenes(ARCHIVO.read_bytes())
    assert len(imagenes) == len(PAGINAS) == 7
    img = Image.open(io.BytesIO(imagenes[0]))
    assert img.format == "JPEG" and max(img.size) == tl.LADO_MAX_PX


# --------------------------------------------------------------- exportación
def test_exportar_excel_formato():
    filas = [
        {"Folio": "folio 18-19", "Fecha": "05/10/2026", "Hora": "06:00", "Transcripción": "aviso 1\naviso 2"},
        {"Folio": "folio 01 (libro nuevo)", "Fecha": "6/10/2026", "Hora": "", "Transcripción": "otro"},
        {"Folio": "folio 02", "Fecha": "fecha rara", "Hora": "25:00", "Transcripción": "x"},
    ]
    ws = load_workbook(io.BytesIO(tl.exportar_excel(filas))).active
    assert [c.value for c in ws[1]] == ["Folio", "Fecha", "Hora", "Transcripción"]
    assert ws.max_column == 4 and ws.max_row == 4
    folio, fecha, hora, texto = ws[2]
    assert folio.value == "folio 18-19" and folio.data_type == "s"
    assert fecha.value == dt.datetime(2026, 10, 5) and fecha.number_format == "DD/MM/YYYY"
    assert hora.value == dt.time(6, 0) and hora.number_format == "HH:MM"
    assert texto.value == "aviso 1\naviso 2" and texto.alignment.wrap_text
    assert ws["A3"].value == "folio 01 (libro nuevo)"
    assert ws["B3"].value == dt.datetime(2026, 10, 6) and ws["C3"].value is None
    assert ws["B4"].value == "fecha rara" and ws["C4"].value == "25:00"  # texto, para corregir


def test_nombre_archivo():
    filas = [{"Fecha": "05/10/2026"}, {"Fecha": ""}, {"Fecha": "7/10/2026"}]
    assert tl.nombre_archivo(filas) == "LG_transcripcion_05-10-2026_07-10-2026.xlsx"
    assert tl.nombre_archivo([{"Fecha": ""}]) == "LG_transcripcion.xlsx"


def test_resumen():
    t = tanda_ok([0], respuesta([fila(texto="SE[?]AL 12 [ilegible] y [?]")]))
    r = tl.resumen(tl.unir_tandas([t, tl.Tanda([1], estado="error")]), [t, tl.Tanda([1], estado="error")])
    assert (r["filas"], r["dudosos"], r["ilegibles"], r["tandas_error"]) == (1, 2, 1, ["pág. 2"])


# ------------------------------------------------------------- configuración
@pytest.fixture
def carpeta(tmp_path):
    shutil.copytree(tl.CARPETA_CONFIG / "originales", tmp_path / "originales")
    return tmp_path


def test_documentos_originales_existen():
    for clave in tl.DOCUMENTOS:
        texto = tl.leer_original(clave)
        assert texto.strip() and tl.tokens_aprox(texto) > 100
    assert "continuacion_del_anterior" in tl.leer_original("instrucciones")


def test_guardar_y_restaurar_documento(carpeta):
    original = tl.leer_documento("referencia", carpeta)
    assert original == tl.leer_original("referencia", carpeta)
    assert not tl.guardar_documento("referencia", original, carpeta)  # sin cambios

    t1 = dt.datetime(2026, 10, 26, 10, 0, 0)
    assert tl.guardar_documento("referencia", "versión 2\r\nnueva", carpeta, ahora=t1)
    assert tl.leer_documento("referencia", carpeta) == "versión 2\nnueva"
    assert tl.guardar_documento("referencia", "versión 3", carpeta, ahora=t1)  # mismo segundo

    versiones = tl.historial("referencia", carpeta)
    assert len(versiones) == 2 and all(f == t1 for f, _ in versiones)
    assert {v.read_text(encoding="utf-8") for _, v in versiones} == {original, "versión 2\nnueva"}
    assert tl.historial("instrucciones", carpeta) == []

    tl.guardar_documento("referencia", tl.leer_original("referencia", carpeta), carpeta)  # restaurar
    assert tl.leer_documento("referencia", carpeta) == original
    assert len(tl.historial("referencia", carpeta)) == 3
    assert tl.leer_original("referencia", carpeta) == original  # el original nunca se toca


def test_ajustes(carpeta):
    assert tl.leer_ajustes(carpeta) == {"modelo": tl.MODELO_DEFECTO}
    tl.guardar_ajustes({"modelo": "claude-sonnet-5-5"}, carpeta)
    assert tl.leer_ajustes(carpeta)["modelo"] == "claude-sonnet-5-5"
    tl.guardar_ajustes({"modelo": "modelo-que-no-existe"}, carpeta)
    assert tl.leer_ajustes(carpeta)["modelo"] == tl.MODELO_DEFECTO


# ------------------------------------------------------- errores y conexión
def _error_api(clase, estado: int, mensaje: str = "error"):
    import anthropic
    import httpx2
    pedido = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
    if clase in (anthropic.APIConnectionError, anthropic.APITimeoutError):
        return clase(request=pedido)
    return clase(mensaje, response=httpx2.Response(estado, request=pedido), body=None)


def test_explicar_error():
    import anthropic
    casos = [
        (_error_api(anthropic.AuthenticationError, 401), "no es válida", True),
        (_error_api(anthropic.NotFoundError, 404), "claude-opus-5-5", True),
        (_error_api(anthropic.BadRequestError, 400, "Your credit balance is too low"), "crédito", True),
        (_error_api(anthropic.RateLimitError, 429), "Esperá", True),
        (_error_api(anthropic.APIConnectionError, 0), "Sistemas", True),
        (_error_api(anthropic.APITimeoutError, 0), "tardó demasiado", False),
        (_error_api(anthropic.InternalServerError, 500), "servidores", False),
        (_error_api(anthropic.BadRequestError, 400, "imagen muy grande"), "imagen muy grande", False),
    ]
    for error, texto, frenar in casos:
        mensaje, frena = tl.explicar_error(error, "claude-opus-5-5")
        assert texto in mensaje and frena is frenar, (type(error).__name__, mensaje)


def test_explicar_error_certificado():
    import anthropic
    e = _error_api(anthropic.APIConnectionError, 0)
    e.__cause__ = Exception("[SSL: CERTIFICATE_VERIFY_FAILED] certificate verify failed")
    mensaje, frena = tl.explicar_error(e, "claude-opus-5-5")
    assert "certificado" in mensaje and frena


def test_probar_conexion():
    cliente = SimpleNamespace(models=SimpleNamespace(
        retrieve=lambda modelo: SimpleNamespace(display_name=f"Nombre de {modelo}")))
    assert tl.probar_conexion(cliente, "claude-opus-5-5") == "Nombre de claude-opus-5-5"
